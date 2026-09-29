"""Tien ich dung file Word giu dinh dang bao cao goc (style Heading2/3, FigCap, TblCap, TableGrid1)."""
import copy, re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


class Doc:
    def __init__(self, template):
        self.d = Document(template)
        body = self.d.element.body
        sect = body[-1]
        for el in list(body):
            if el is not sect:
                body.remove(el)
        self.body = body
        self.sect = sect
        self.bm = 900

    # ------------------------------------------------------------------ doan van
    def _p(self, style=None):
        p = self.d.add_paragraph(style=style)
        return p

    def _runs(self, p, text, size=None, bold=None):
        """Mini markup: _{..} chi so duoi, ^{..} chi so tren, **..** dam, *..* nghieng."""
        text = text.replace('⁻¹', '^{−1}').replace('⁶', '^{6}')
        tokens = re.split(r'(_\{[^}]*\}|\^\{[^}]*\}|\*\*[^*]+\*\*|\*[^*]+\*)', text)
        for tk in tokens:
            if not tk:
                continue
            if tk.startswith('_{'):
                r = p.add_run(tk[2:-1]); r.font.subscript = True
            elif tk.startswith('^{'):
                r = p.add_run(tk[2:-1]); r.font.superscript = True
            elif tk.startswith('**'):
                r = p.add_run(tk[2:-2]); r.bold = True
            elif tk.startswith('*'):
                r = p.add_run(tk[1:-1]); r.italic = True
            else:
                r = p.add_run(tk)
            if size:
                r.font.size = Pt(size)
            if bold is not None and not tk.startswith('**'):
                r.bold = bold
        return p

    def para(self, text, indent=True):
        p = self._p()
        pf = p.paragraph_format
        if indent:
            pf.first_line_indent = Cm(1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._runs(p, text)
        return p

    def lead(self, head, text):
        """Doan co cum mo dau in nghieng dam (a), b)...)."""
        p = self._p(); p.paragraph_format.first_line_indent = Cm(1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(head + ' '); r.bold = True; r.italic = True
        self._runs(p, text)
        return p

    def heading(self, text, level):
        p = self._p('Heading %d' % level)
        self._bookmark(p)
        p.add_run(text)
        return p

    def _bookmark(self, p):
        self.bm += 1
        s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(self.bm)); s.set(qn('w:name'), '_TocN%d' % self.bm)
        e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(self.bm))
        p._p.append(s); p._p.append(e)

    # ------------------------------------------------------------------ hinh
    def figure(self, path, caption, width_cm=15.0):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format; pf.keep_with_next = True; pf.first_line_indent = Cm(0)
        pf.space_before = Pt(5); pf.space_after = Pt(3)
        p.add_run().add_picture(path, width=Cm(width_cm))
        c = self._p('FigCap'); self._bookmark(c); self._runs(c, caption)
        return c

    # ------------------------------------------------------------------ bang
    def table(self, caption, header, rows, widths, size=11, bold_rows=(), note=None, header_rows=1,
              align=None, merge=()):
        c = self._p('TblCap'); self._bookmark(c); self._runs(c, caption)
        c.paragraph_format.keep_with_next = True
        ncol = len(widths)
        allrows = header + rows
        t = self.d.add_table(rows=len(allrows), cols=ncol)
        tbl = t._tbl
        tblPr = tbl.tblPr
        st = OxmlElement('w:tblStyle'); st.set(qn('w:val'), 'TableGrid1')
        tblPr.insert(0, st)
        for old in tblPr.findall(qn('w:tblStyle'))[1:]:
            tblPr.remove(old)
        tw = tblPr.find(qn('w:tblW'))
        if tw is None:
            tw = OxmlElement('w:tblW'); tblPr.insert(1, tw)
        tw.set(qn('w:w'), '9072'); tw.set(qn('w:type'), 'dxa')
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tw.addnext(jc)
        tot = sum(widths)
        grid = tbl.tblGrid
        tw_total = 9072  # 16 cm
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(round(w / tot * tw_total))))
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed')
        tblPr.find(qn('w:jc')).addnext(lay)
        for i, row in enumerate(allrows):
            tr = t.rows[i]
            if i < header_rows:
                trPr = tr._tr.get_or_add_trPr()
                h = OxmlElement('w:tblHeader'); trPr.append(h)
            for j in range(ncol):
                cell = tr.cells[j]
                cell.width = Cm(widths[j] / tot * 16.0)
                txt = row[j] if j < len(row) else ''
                p = cell.paragraphs[0]
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0); pf.space_before = Pt(1); pf.space_after = Pt(1)
                pf.line_spacing = 1.0
                if i < header_rows:
                    pf.keep_with_next = True
                a = 'center'
                if align:
                    a = align[j]
                p.alignment = {'center': WD_ALIGN_PARAGRAPH.CENTER, 'left': WD_ALIGN_PARAGRAPH.LEFT}[a]
                b = (i < header_rows) or (i - header_rows) in bold_rows
                self._runs(p, txt, size=size, bold=b if b else None)
                tcPr = cell._tc.get_or_add_tcPr()
                va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
        for (r0, c0, r1, c1) in merge:
            a = t.cell(r0, c0); b = t.cell(r1, c1)
            txt = a.text
            m = a.merge(b)
            # giu mot doan van
            for p in m.paragraphs[1:]:
                p._p.getparent().remove(p._p)
        if note:
            p = self._p(); p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(2)
            self._runs(p, note, size=11)
            for r in p.runs:
                r.italic = True
        return t

    # ------------------------------------------------------------------ cong thuc
    def equation(self, omml_inner, number):
        p = self._p()
        pPr = p._p.get_or_add_pPr()
        tabs = parse_xml('<w:tabs xmlns:w="%s"><w:tab w:val="center" w:pos="4536"/><w:tab w:val="right" w:pos="9071"/></w:tabs>' % W_NS)
        pPr.append(tabs)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
        p.add_run('\t')
        om = parse_xml('<m:oMath xmlns:m="%s" xmlns:w="%s">%s</m:oMath>' % (M_NS, W_NS, omml_inner))
        p._p.append(om)
        r = p.add_run('\t(%s)' % number)
        return p

    def save(self, path):
        # bo cac quan he anh cua noi dung goc khong con dung
        xml = self.d.element.xml
        part = self.d.part
        for rid, rel in list(part.rels.items()):
            if rel.reltype.endswith('/image') and ('"%s"' % rid) not in xml:
                del part.rels[rid]
        self.d.save(path)


# ---------------------------------------------------------------------- OMML nho gon
def mr(t, plain=True):
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ''
    t = t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    return '%s<m:r>%s<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t xml:space="preserve">%s</m:t></m:r>' % ('', sty, t)


def mi(t):
    return mr(t, plain=False)


def sub(base, s):
    return '<m:sSub><m:e>%s</m:e><m:sub>%s</m:sub></m:sSub>' % (base, s)


def sup(base, s):
    return '<m:sSup><m:e>%s</m:e><m:sup>%s</m:sup></m:sSup>' % (base, s)


def frac(n, d):
    return '<m:f><m:num>%s</m:num><m:den>%s</m:den></m:f>' % (n, d)


def paren(x, beg='(', end=')'):
    return '<m:d><m:dPr><m:begChr m:val="%s"/><m:endChr m:val="%s"/></m:dPr><m:e>%s</m:e></m:d>' % (beg, end, x)


def sqrt(x):
    return '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>%s</m:e></m:rad>' % x


def nary(chr_, lo, hi, body):
    return ('<m:nary><m:naryPr><m:chr m:val="%s"/></m:naryPr><m:sub>%s</m:sub><m:sup>%s</m:sup><m:e>%s</m:e></m:nary>'
            % (chr_, lo, hi, body))
