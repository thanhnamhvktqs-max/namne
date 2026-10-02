# -*- coding: utf-8 -*-
"""Bo cong cu dung slide: hinh khoi, chu nhieu dinh dang, mui ten, chuyen canh,
va bo sinh XML hieu ung (p:timing) theo dung cau truc PowerPoint."""
import itertools
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from data import INK

FONT = "Arial"
MATH = "Cambria"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def E(x):
    return Emu(int(round(x * 914400)))


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------
# Chu nhieu dinh dang: **dam**, *nghieng*, _{chi so duoi}, ^{chi so tren},
# [[RRGGBB:doan doi mau]], ~~...~~ (font cong thuc Cambria)
# ---------------------------------------------------------------------------
_TOK = re.compile(r"(\*\*|\*|_\{|\^\{|\}|\[\[[0-9A-Fa-f]{6}:|\]\]|~~)")


def parse_markup(s):
    runs = []
    bold = italic = math = False
    stack = []          # cac nhom { } dang mo: 'sub' | 'sup'
    colors = []
    for part in _TOK.split(s):
        if part == "":
            continue
        if part == "**":
            bold = not bold
        elif part == "*":
            italic = not italic
        elif part == "~~":
            math = not math
        elif part == "_{":
            stack.append("sub")
        elif part == "^{":
            stack.append("sup")
        elif part == "}" and stack:
            stack.pop()
        elif part.startswith("[[") and part.endswith(":"):
            colors.append(part[2:8])
        elif part == "]]" and colors:
            colors.pop()
        else:
            base = None
            if stack:
                base = -25000 if stack[-1] == "sub" else 30000
            runs.append(dict(t=part, b=bold, i=italic, base=base,
                             color=colors[-1] if colors else None, math=math))
    return runs


def _apply_rpr(r, size, bold, italic, color, font, base=None, spc=None):
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rPr = r._r.get_or_add_rPr()
    rPr.set("lang", "vi-VN")
    if base:
        rPr.set("baseline", str(base))
    if spc:
        rPr.set("spc", str(spc))
    # font cho chu dong A va chu phuc hop de chu Viet khong bi thay font
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def fill_text(tf, content, size=18, color=INK, bold=False, italic=False, font=FONT,
              align="l", line=None, space_after=0, space_before=0, spc=None, bullet=None):
    """content: chuoi (\n tach doan) hoac list doan; moi doan la chuoi hoac dict
    {t, size, color, bold, italic, align, bullet, space_after, space_before, line, font}"""
    if isinstance(content, str):
        paras = content.split("\n")
    else:
        paras = content
    first = True
    for para in paras:
        spec = para if isinstance(para, dict) else {"t": para}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        al = spec.get("align", align)
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT,
                       "j": PP_ALIGN.JUSTIFY}[al]
        ln = spec.get("line", line)
        if ln:
            p.line_spacing = ln
        sa = spec.get("space_after", space_after)
        sb = spec.get("space_before", space_before)
        if sa:
            p.space_after = Pt(sa)
        if sb:
            p.space_before = Pt(sb)
        sz = spec.get("size", size)
        bl = spec.get("bullet", bullet)
        if bl:
            pPr = p._p.get_or_add_pPr()
            ind = int(spec.get("indent", 0.22) * 914400)
            pPr.set("marL", str(ind))
            pPr.set("indent", str(-ind))
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", bl if isinstance(bl, str) and len(bl) == 6 else "0A5FA8")
            etree.SubElement(pPr, qn("a:buSzPct")).set("val", "100000")
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", spec.get("char", "•"))
        for run in parse_markup(spec.get("t", "")):
            r = p.add_run()
            r.text = run["t"]
            _apply_rpr(r, sz, bool(spec.get("bold", bold)) or run["b"],
                       bool(spec.get("italic", italic)) or run["i"],
                       run["color"] or spec.get("color", color),
                       MATH if run["math"] else spec.get("font", font), run["base"],
                       spec.get("spc", spc))
    return tf


def _no_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _set_alpha(shape, alpha):
    """alpha 0..100 (% do dac) cho mau to dac."""
    sf = shape._element.spPr.find(qn("a:solidFill"))
    if sf is not None:
        clr = sf[0]
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))


def shadow(shape, blur=0.10, dist=0.035, alpha=16, color="0B2545"):
    spPr = shape._element.spPr
    old = spPr.find(qn("a:effectLst"))
    if old is not None:
        spPr.remove(old)
    eff = etree.Element(qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 914400)),
                          dist=str(int(dist * 914400)), dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 1000)))
    # effectLst dung sau a:ln (hoac sau phan to mau) trong spPr
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(eff)
    else:
        anchor = None
        for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:prstGeom", "a:custGeom"):
            anchor = spPr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addnext(eff)
        else:
            spPr.append(eff)


def box(S, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, kind=None, alpha=None,
        dash=None, shd=False, name=None, text=None, rot=0, lsp=None, **tk):
    # 'line' la mau duong vien; neu truyen so thi hieu la gian dong cua chu (tuong thich text())
    if isinstance(line, (int, float)) and not isinstance(line, bool):
        lsp, line = line, None
    if lsp is not None:
        tk["line"] = lsp
    if kind is None:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = S.add_shape(kind, E(x), E(y), E(w), E(h))
    _no_style(sh)
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _set_alpha(sh, alpha)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    else:
        sh.line.fill.background()
    if rot:
        sh.rotation = rot
    if shd:
        shadow(sh)
    if name:
        sh.name = name
    tf = sh.text_frame
    m = tk.pop("margin", 0.06)
    if isinstance(m, (int, float)):
        m = (m, m, m, m)
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in m]
    tf.word_wrap = tk.pop("wrap", True)
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[tk.pop("anchor", "m")]
    if text is not None:
        tk.setdefault("align", "c")
        fill_text(tf, text, **tk)
    return sh


def text(S, x, y, w, h, content, size=18, anchor="t", margin=0.0, wrap=True, name=None, rot=0, **tk):
    tb = S.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    if isinstance(margin, (int, float)):
        margin = (margin, margin, margin, margin)
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in margin]
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_text(tf, content, size=size, **tk)
    if name:
        tb.name = name
    if rot:
        tb.rotation = rot
    return tb


def line(S, x1, y1, x2, y2, color=INK, w=1.5, head=None, tail="tri", dash=None, name=None, size="med"):
    c = S.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _no_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    if head:
        etree.SubElement(ln, qn("a:headEnd"), type={"tri": "triangle", "oval": "oval"}[head], w=size, len=size)
    if tail:
        etree.SubElement(ln, qn("a:tailEnd"), type={"tri": "triangle", "oval": "oval"}[tail], w=size, len=size)
    if name:
        c.name = name
    return c


def pic(S, path, x, y, w=None, h=None, crop=None, name=None):
    """crop = (trai, tren, phai, duoi) theo ti le 0..1. Neu chi co w hoac h -> giu ti le."""
    from PIL import Image
    iw, ih = Image.open(path).size
    if crop:
        cl, ct, cr, cb = crop
        iw2, ih2 = iw * (1 - cl - cr), ih * (1 - ct - cb)
    else:
        iw2, ih2 = iw, ih
    if w is None and h is None:
        w = iw2 / 150.0
    if h is None:
        h = w * ih2 / iw2
    if w is None:
        w = h * iw2 / ih2
    p = S.add_picture(path, E(x), E(y), E(w), E(h))
    if crop:
        p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = crop
    if name:
        p.name = name
    return p


def icon_disc(S, cx, cy, d, icon_path, fill, ring=None, name=None, icon_scale=0.56):
    """Hinh tron mau + bieu tuong trang, gom thanh 1 nhom de lam hieu ung."""
    g = S.add_group_shape()
    box(g.shapes, cx - d / 2, cy - d / 2, d, d, fill=fill, kind=MSO_SHAPE.OVAL,
        line=ring, lw=2.0 if ring else 1.0)
    s = d * icon_scale
    pic(g.shapes, icon_path, cx - s / 2, cy - s / 2, s, s)
    if name:
        g.name = name
    return g


def slide_number(slide, n, x=12.33, y=7.04, color="5B6B7B"):
    tb = text(slide.shapes, x, y, 0.6, 0.3, "", size=12, align="r")
    p = tb.text_frame.paragraphs[0]._p
    fld = etree.SubElement(p, qn("a:fld"), id="{8D2B4A71-3C1E-4C55-9E2B-0A1B2C3D4E5F}", type="slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"), lang="vi-VN", sz="1200")
    sf = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(rpr, qn("a:latin"), typeface=FONT)
    t = etree.SubElement(fld, qn("a:t"))
    t.text = str(n)
    tb.name = "SlideNo"
    return tb


# ---------------------------------------------------------------------------
# Chuyen canh
# ---------------------------------------------------------------------------
def set_transition(slide, kind="fade", spd="med", **kw):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    for old in sld.findall("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
        sld.remove(old)
    if kind == "morph":
        xml = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
               'xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
               f'xmlns:p="{P_NS}"><mc:Choice Requires="p159"><p:transition spd="slow">'
               '<p159:morph option="byObject"/></p:transition></mc:Choice><mc:Fallback>'
               '<p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    else:
        inner = {
            "fade": "<p:fade/>",
            "push": f'<p:push dir="{kw.get("dir", "u")}"/>',
            "wipe": f'<p:wipe dir="{kw.get("dir", "r")}"/>',
            "zoom": '<p:zoom dir="in"/>',
            "cover": f'<p:cover dir="{kw.get("dir", "l")}"/>',
            "split": '<p:split orient="vert" dir="out"/>',
        }[kind]
        xml = f'<p:transition xmlns:p="{P_NS}" spd="{spd}">{inner}</p:transition>'
    el = etree.fromstring(xml)
    clr = sld.find(qn("p:clrMapOvr"))
    clr.addnext(el)


# ---------------------------------------------------------------------------
# Hieu ung (p:timing)
# ---------------------------------------------------------------------------
EFFECTS = {
    # ten: (presetID, class, subtype, filter)
    "appear": (1, "entr", 0, None),
    "fade": (10, "entr", 0, "fade"),
    "wipe_l": (22, "entr", 8, "wipe(left)"),
    "wipe_r": (22, "entr", 2, "wipe(right)"),
    "wipe_t": (22, "entr", 1, "wipe(up)"),
    "wipe_b": (22, "entr", 4, "wipe(down)"),
    "split": (16, "entr", 42, "barn(outVertical)"),
    "zoom": (53, "entr", 16, None),
    "rise": (42, "entr", 0, None),
    "fly_b": (2, "entr", 4, None),
    "fly_l": (2, "entr", 8, None),
    "pulse": (6, "emph", 0, None),
    "exit_fade": (10, "exit", 0, "fade"),
    "path": (0, "path", 0, None),
}


class Anim:
    """Danh sach nhom hieu ung theo lan nhap.

    a = Anim(slide); a.auto() hoac a.click(); a.add(shape, 'fade'); a.add(s2, 'zoom', after=True)
    - Phan tu dau cua nhom: Click (hoac After Previous neu nhom auto dau slide)
    - after=False: With Previous (cong them delay so voi hieu ung lien truoc)
    - after=True : After Previous (bat dau khi nhom con truoc ket thuc)
    """

    def __init__(self, slide):
        self.slide = slide
        self.groups = []
        self.grp = {}
        self.bld = []

    def click(self):
        self.groups.append({"trigger": "click", "pars": []})
        return self

    def auto(self):
        assert not self.groups, "nhom auto chi dung o dau slide"
        self.groups.append({"trigger": "auto", "pars": []})
        return self

    def add(self, shape, eff, dur=500, delay=0, after=False, **kw):
        if not self.groups:
            self.click()
        g = self.groups[-1]
        spid = shape.shape_id
        gid = self.grp.get(spid, -1) + 1
        self.grp[spid] = gid
        if shape._element.tag == qn("p:sp"):
            self.bld.append((spid, gid))
        total = dur * 2 if eff == "pulse" else dur
        if not g["pars"]:
            g["pars"].append({"start": 0, "effects": []})
            node = "clickEffect" if g["trigger"] == "click" else "afterEffect"
            d = delay
        elif after:
            par = g["pars"][-1]
            end = max(e["delay"] + e["total"] for e in par["effects"])
            g["pars"].append({"start": par["start"] + end, "effects": []})
            node, d = "afterEffect", delay
        else:
            par = g["pars"][-1]
            node, d = "withEffect", par["effects"][-1]["delay"] + delay
        g["pars"][-1]["effects"].append(dict(spid=spid, eff=eff, dur=dur, delay=d, node=node,
                                             gid=gid, total=total, kw=kw))
        return self

    # -- sinh XML --------------------------------------------------------------
    def _eff(self, e, ids):
        pid, cls, sub, filt = EFFECTS[e["eff"]]
        spid, dur = e["spid"], e["dur"]
        tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
        eid = next(ids)
        extra = ""
        if e["eff"] == "fly_b" or e["eff"] == "fly_l":
            extra = 'decel="60000" '
        if e["eff"] == "path" and e["kw"].get("ease", True):
            extra = 'accel="30000" decel="30000" '
        head = (f'<p:par><p:cTn id="{eid}" presetID="{pid}" presetClass="{cls}" presetSubtype="{sub}" '
                f'{extra}fill="hold" grpId="{e["gid"]}" nodeType="{e["node"]}"><p:stCondLst>'
                f'<p:cond delay="{e["delay"]}"/></p:stCondLst><p:childTnLst>')
        tail = "</p:childTnLst></p:cTn></p:par>"

        def setvis(val="visible", at=0):
            i = next(ids)
            return (f'<p:set><p:cBhvr><p:cTn id="{i}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{at}"/>'
                    f'</p:stCondLst></p:cTn>{tgt}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                    f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="{val}"/></p:to></p:set>')

        def effect(f, tr="in"):
            i = next(ids)
            return (f'<p:animEffect transition="{tr}" filter="{f}"><p:cBhvr><p:cTn id="{i}" dur="{dur}"/>'
                    f'{tgt}</p:cBhvr></p:animEffect>')

        def prop(attr, v0, v1):
            i = next(ids)
            return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{i}" dur="{dur}" fill="hold"/>'
                    f'{tgt}<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                    f'<p:tav tm="0"><p:val><p:strVal val="{v0}"/></p:val></p:tav><p:tav tm="100000"><p:val>'
                    f'<p:strVal val="{v1}"/></p:val></p:tav></p:tavLst></p:anim>')

        x = e["eff"]
        if x == "appear":
            body = setvis()
        elif x in ("fade", "wipe_l", "wipe_r", "wipe_t", "wipe_b", "split"):
            body = setvis() + effect(filt)
        elif x == "zoom":
            body = setvis() + effect("fade") + prop("ppt_w", "0", "#ppt_w") + prop("ppt_h", "0", "#ppt_h")
        elif x == "rise":
            body = setvis() + effect("fade") + prop("ppt_x", "#ppt_x", "#ppt_x") + prop("ppt_y", "#ppt_y+.05", "#ppt_y")
        elif x == "fly_b":
            body = setvis() + prop("ppt_x", "#ppt_x", "#ppt_x") + prop("ppt_y", "1+#ppt_h/2", "#ppt_y")
        elif x == "fly_l":
            body = setvis() + prop("ppt_x", "0-#ppt_w/2", "#ppt_x") + prop("ppt_y", "#ppt_y", "#ppt_y")
        elif x == "pulse":
            i = next(ids)
            sc = e["kw"].get("scale", 112)
            body = (f'<p:animScale><p:cBhvr><p:cTn id="{i}" dur="{dur}" autoRev="1" fill="hold"/>{tgt}'
                    f'</p:cBhvr><p:by x="{sc * 1000}" y="{sc * 1000}"/></p:animScale>')
        elif x == "exit_fade":
            body = effect("fade", "out") + setvis("hidden", max(dur - 1, 0))
        elif x == "path":
            i = next(ids)
            pth = e["kw"]["path"]
            npts = pth.count(" L ") + 1
            body = (f'<p:animMotion origin="layout" path="{pth}" pathEditMode="relative" rAng="0" '
                    f'ptsTypes="{"A" * (npts + 0)}"><p:cBhvr><p:cTn id="{i}" dur="{dur}" fill="hold"/>{tgt}'
                    f'<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName>'
                    f'</p:attrNameLst></p:cBhvr><p:rCtr x="0" y="0"/></p:animMotion>')
        else:
            raise ValueError(x)
        return head + body + tail

    def build(self):
        if not self.groups:
            return
        ids = itertools.count(3)
        groups = []
        for g in self.groups:
            if not g["pars"]:
                continue
            oid = next(ids)
            cond = '<p:cond delay="indefinite"/>'
            if g["trigger"] == "auto":
                cond += '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
            inner = []
            for par in g["pars"]:
                pid = next(ids)
                effs = "".join(self._eff(e, ids) for e in par["effects"])
                inner.append(f'<p:par><p:cTn id="{pid}" fill="hold"><p:stCondLst><p:cond delay="{par["start"]}"/>'
                             f'</p:stCondLst><p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par>')
            groups.append(f'<p:par><p:cTn id="{oid}" fill="hold"><p:stCondLst>{cond}</p:stCondLst>'
                          f'<p:childTnLst>{"".join(inner)}</p:childTnLst></p:cTn></p:par>')
        bld = "".join(f'<p:bldP spid="{s}" grpId="{g}" animBg="1"/>' for s, g in self.bld)
        xml = (f'<p:timing xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" '
               'restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
               f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(groups)}'
               '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/>'
               '</p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl>'
               '<p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par>'
               f'</p:tnLst>{"<p:bldLst>" + bld + "</p:bldLst>" if bld else ""}</p:timing>')
        sld = self.slide._element
        for old in sld.findall(qn("p:timing")):
            sld.remove(old)
        el = etree.fromstring(xml)
        ext = sld.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)

    def clicks(self):
        return sum(1 for g in self.groups if g["trigger"] == "click" and g["pars"])


__all__ = ["E", "rgb", "box", "text", "line", "pic", "icon_disc", "shadow", "slide_number",
           "set_transition", "Anim", "fill_text", "MSO_SHAPE", "MSO_LINE_DASH_STYLE", "FONT", "MATH"]
