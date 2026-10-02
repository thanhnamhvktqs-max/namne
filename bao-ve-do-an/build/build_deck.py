# -*- coding: utf-8 -*-
"""Dung bo slide bao ve do an (30 slide, 16:9) tu bo slide goc.

    python build_deck.py <slide_goc.pptx> <thu_muc_assets> <file_ra.pptx>

thu_muc_assets chua: img/ (prep_assets.py), charts/ (charts.py), icons/ (icon trang).
Nen (luc giac xanh, song xanh) va slide master duoc giu nguyen tu bo slide goc.
"""
import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from data import (ACCENT, BAD, BEFORE, BLUE, CARD, GOOD, INK, MUTED, NAVY, PITCH, TINT,
                  TRADE, YAW, STEPS)
from kit import (Anim, E, MSO_LINE_DASH_STYLE, MSO_SHAPE, box, icon_disc, line, pic,
                 set_transition, slide_number, text)

DASH = MSO_LINE_DASH_STYLE.DASH
from notes import NOTES

SRC, ASSETS, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
IMG = os.path.join(ASSETS, "img")
CH = os.path.join(ASSETS, "charts")
ICO = os.path.join(ASSETS, "icons")
ANCH = json.load(open(os.path.join(CH, "anchors.json"), encoding="utf-8"))

LM, RM = 0.55, 12.80
CW = RM - LM
ACC_TINT = "FFF1E6"
BAD_TINT = "FDECEC"
GOOD_TINT = "E8F5EC"
LINE_SOFT = "D5DEE8"

prs = Presentation(SRC)
# --- xoa toan bo slide cu (giu master, layout, nen) ---
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
# --- font chu de -> Arial (moi chu khong dat font rieng cung dung Arial) ---
for rel in prs.slide_master.part.rels.values():
    if rel.reltype.endswith("/theme"):
        th = rel.target_part
        from lxml import etree
        root = etree.fromstring(th.blob)
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for tag in ("majorFont", "minorFont"):
            el = root.find(f".//{A}{tag}/{A}latin")
            if el is not None:
                el.set("typeface", "Arial")
        th._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

LAYOUT = {}
for lay in prs.slide_layouts:
    LAYOUT[lay.name] = lay
L_HEX = LAYOUT["Divider Slide 2D"]
L_WAVE = LAYOUT["Divider Slide 1E"]

_icon_cache = {}


def icon(name, color="FFFFFF"):
    """Duong dan icon PNG theo mau (to lai mau tu ban trang)."""
    if color.upper() == "FFFFFF":
        p = os.path.join(ICO, name + ".png")
        if os.path.exists(p):
            return p
        return os.path.join(IMG, name + ".png")
    key = (name, color)
    if key in _icon_cache:
        return _icon_cache[key]
    src = Image.open(os.path.join(ICO, name + ".png")).convert("RGBA")
    r, g, b = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    col = Image.new("RGBA", src.size, (r, g, b, 255))
    col.putalpha(src.split()[3])
    p = os.path.join(ICO, f"{name}_{color}.png")
    col.save(p)
    _icon_cache[key] = p
    return p


def img(name):
    for ext in (".png", ".jpg"):
        p = os.path.join(IMG, name + ext)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(name)


def chart(name, layer="base"):
    return os.path.join(CH, ANCH[name]["files"][layer])


SLIDES = []


def new_slide(layout=None):
    s = prs.slides.add_slide(layout or L_HEX)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    SLIDES.append(s)
    return s


def header(s, section, title, size=27):
    S = s.shapes
    text(S, LM, 0.22, 10.0, 0.32, section, size=13, bold=True, color=BLUE, spc=60, name="Section")
    return text(S, LM, 0.52, CW, 0.66, title, size=size, bold=True, color=NAVY, anchor="m",
                name="Title")


def card(S, x, y, w, h, fill=CARD, shd=True, line_=None, radius=0.10, alpha=None, name=None, lw=1.0):
    return box(S, x, y, w, h, fill=fill, radius=radius, shd=shd, line=line_, lw=lw, alpha=alpha, name=name)


def label(S, x, y, w, txt, color=BLUE, ico=None, size=12.5, h=0.32):
    """Nhan nho in hoa co bieu tuong (bieu tuong cung mau)."""
    off = 0
    if ico:
        pic(S, icon(ico, color), x, y + (h - 0.26) / 2, 0.26, 0.26)
        off = 0.36
    return text(S, x + off, y, w - off, h, txt, size=size, bold=True, color=color, anchor="m", spc=40)


def group(S):
    return S.add_group_shape()


def kpi(S, x, y, w, h, value, lab, color=ACCENT, vsize=26, lsize=12.5, fill=CARD, name=None):
    g = group(S)
    card(g.shapes, x, y, w, h, fill=fill)
    text(g.shapes, x + 0.08, y + 0.05, w - 0.16, h * 0.56, value, size=vsize, bold=True, color=color,
         align="c", anchor="m")
    text(g.shapes, x + 0.08, y + h * 0.58, w - 0.16, h * 0.40, lab, size=lsize, color=INK, align="c",
         anchor="t", line=0.95)
    if name:
        g.name = name
    return g


def chip(S, x, y, w, h, txt, fill=NAVY, color="FFFFFF", size=13, bold=True, radius=None, line_=None, **kw):
    return box(S, x, y, w, h, fill=fill, line=line_, radius=radius if radius is not None else h / 2,
               text=txt, size=size, bold=bold, color=color, margin=(0.08, 0.02, 0.08, 0.02), **kw)


def arrow(S, x1, y1, x2, y2, color="7F8B98", w=1.75, **kw):
    return line(S, x1, y1, x2, y2, color=color, w=w, **kw)


def table(S, x, y, widths, rh, rows, size=14, hl=None, head_fill=NAVY, aligns=None, bold_cols=(),
          zebra=True, row_h=None):
    """Bang tu ve bang hinh khoi: tra ve danh sach nhom (moi hang 1 nhom) de lam hieu ung."""
    out = []
    cy = y
    for ri, row in enumerate(rows):
        h = (row_h[ri] if row_h else rh)
        g = group(S)
        if ri == 0:
            fill, col, bold = head_fill, "FFFFFF", True
        elif hl is not None and ri == hl:
            fill, col, bold = ACC_TINT, INK, True
        else:
            fill = ("F4F8FC" if ri % 2 == 0 else "FFFFFF") if zebra else "FFFFFF"
            col, bold = INK, False
        box(g.shapes, x, cy, sum(widths), h, fill=fill, line=(ACCENT if (hl is not None and ri == hl) else None),
            lw=1.5)
        cx = x
        for ci, cell in enumerate(row):
            al = (aligns[ci] if aligns else ("l" if ci == 0 else "c"))
            text(g.shapes, cx + 0.06, cy, widths[ci] - 0.12, h, cell, size=size if ri else size - 0.5,
                 bold=bold or (ci in bold_cols and ri > 0), color=col, align=al, anchor="m", line=0.95)
            cx += widths[ci]
        out.append(g)
        cy += h
    return out


def number(s):
    slide_number(s, len(SLIDES))


def notes(s, key):
    s.notes_slide.notes_text_frame.text = NOTES[key]


# =============================================================================
# 1. TRANG BIA
# =============================================================================
def s01():
    s = new_slide()
    S = s.shapes
    logo = pic(S, img("logo"), 0.55, 0.38, 1.12, 1.12)
    inst = text(S, 1.85, 0.42, 7.0, 1.05, [
        {"t": "HỌC VIỆN KỸ THUẬT QUÂN SỰ", "size": 21, "bold": True, "color": "B71C1C"},
        {"t": "VIỆN TÊN LỬA VÀ KỸ THUẬT ĐIỀU KHIỂN", "size": 16, "bold": True, "color": BLUE},
    ], anchor="m", line=1.1)
    tag = chip(S, LM, 1.88, 4.3, 0.42, "BÁO CÁO ĐỒ ÁN TỐT NGHIỆP 2026", fill=None, color=NAVY,
               size=14, line_=NAVY, radius=0.21, spc=60)
    title = text(S, LM, 2.48, 8.1, 1.62, "THIẾT KẾ BỘ ĐIỀU KHIỂN GIMBAL\nTRÊN CƠ SỞ HỆ ĐIỀU HÀNH\nTHỜI GIAN THỰC NUTTX",
                 size=31, bold=True, color=NAVY, line=1.02, anchor="m")
    hero = pic(S, img("cad_render_cut"), 8.75, 1.05, 4.15)
    g = group(S)
    box(g.shapes, LM, 4.45, 8.0, 1.85, fill="FFFFFF", alpha=72, radius=0.12)
    text(g.shapes, LM + 0.2, 4.6, 1.85, 0.62, "Giảng viên\nhướng dẫn", size=14, color=MUTED, line=0.95)
    text(g.shapes, LM + 2.0, 4.58, 5.95, 0.72, [
        "Trung tá, TS. **Nguyễn Ngọc Hưng** – Viện Tên lửa và KTĐK",
        "Trung tá, ThS. **Lê Văn Huy** – Trường Cao đẳng Kỹ thuật PK-KQ"], size=14, color=INK, line=1.08)
    text(g.shapes, LM + 0.2, 5.5, 1.85, 0.62, "Học viên\nthực hiện", size=14, color=MUTED, line=0.95)
    text(g.shapes, LM + 2.0, 5.55, 5.95, 0.4, "**Huỳnh Thanh Nam** – Lớp Tên lửa phòng không 2, c157, d1",
         size=14, color=INK, anchor="m")
    date = text(S, 8.75, 5.75, 4.05, 0.4, "Hà Nội, tháng 10 năm 2026", size=15, italic=True, color=NAVY,
                align="c")
    a = Anim(s)
    a.auto().add(hero, "zoom", 700)
    a.add(title, "wipe_l", 600, after=True)
    a.add(tag, "fade", 400, delay=150)
    a.add(g, "rise", 500, after=True)
    a.add(date, "fade", 400, delay=200)
    a.build()
    set_transition(s, "fade", "slow")
    number(s)
    notes(s, 1)


# =============================================================================
# 2. LY DO CHON DE TAI - UNG DUNG UAV CANH BANG
# =============================================================================
def uav_scene(S, ox, oy):
    """Ve UAV canh bang (nhin ngang) bang hinh khoi ban dia. Tra ve cac nhom."""
    g = group(S)
    G = g.shapes
    body = "4B5D73"
    dark = "2F3E52"
    # than
    box(G, ox + 0.0, oy + 0.30, 3.05, 0.44, fill=body, radius=0.22)
    # mui (bo tron them)
    box(G, ox + 2.75, oy + 0.33, 0.45, 0.38, fill=body, kind=MSO_SHAPE.OVAL)
    # canh (mat cat ngang)
    box(G, ox + 1.05, oy + 0.24, 1.25, 0.16, fill=dark, kind=MSO_SHAPE.OVAL)
    # duoi dung
    box(G, ox + 0.02, oy - 0.30, 0.55, 0.62, fill=body, kind=MSO_SHAPE.RIGHT_TRIANGLE)
    # duoi ngang
    box(G, ox - 0.05, oy + 0.34, 0.75, 0.11, fill=dark, kind=MSO_SHAPE.OVAL)
    # canh quat day phia sau
    box(G, ox - 0.22, oy + 0.08, 0.09, 0.88, fill="9AA7B4", kind=MSO_SHAPE.OVAL)
    box(G, ox - 0.16, oy + 0.45, 0.16, 0.14, fill=dark, kind=MSO_SHAPE.OVAL)
    # tum gimbal duoi mui
    box(G, ox + 2.05, oy + 0.62, 0.46, 0.46, fill="1E2B3A", kind=MSO_SHAPE.OVAL)
    box(G, ox + 2.21, oy + 0.80, 0.17, 0.17, fill="4FC3F7", kind=MSO_SHAPE.OVAL)
    return g


def s02():
    s = new_slide()
    S = s.shapes
    header(s, "MỞ ĐẦU · LÝ DO CHỌN ĐỀ TÀI", "Camera trên UAV cánh bằng cần được ổn định đường ngắm")
    # --- khung canh minh hoa ---
    card(S, LM, 1.38, 6.55, 5.24)
    text(S, LM + 0.2, 1.48, 6.0, 0.32, "UAV CÁNH BẰNG MANG CAMERA QUAN SÁT", size=12.5, bold=True,
         color=BLUE, spc=40)
    ground = box(S, LM + 0.12, 5.95, 6.31, 0.55, fill="DCE8DF", radius=0.08)
    tx, ty = 5.85, 5.92  # muc tieu
    target = icon_disc(S, tx, ty, 0.5, icon("crosshairs"), BAD)
    tlab = text(S, tx - 0.85, 6.2, 1.7, 0.3, "mục tiêu", size=12.5, color=INK, align="c")
    uav = uav_scene(S, 1.05, 2.55)
    tcx, tcy = 1.05 + 2.28, 2.55 + 0.85   # tam tum gimbal
    # duong ngam bi lech (gan cung)
    bads = []
    for dx in (-1.25, 0.95, -0.35):
        bads.append(line(S, tcx, tcy, tx + dx, ty, color=BAD, w=2.0, tail=None, dash=DASH))
    bad_lab = text(S, 1.0, 4.95, 2.7, 0.62, "Gắn cứng: đường ngắm bị\nkéo lệch, ảnh rung, mất mục tiêu",
                   size=13, color=BAD, bold=True, line=1.0)
    # nhieu
    d1 = group(S)
    pic(d1.shapes, icon("wind", MUTED), 2.2, 1.95, 0.34, 0.34)
    text(d1.shapes, 2.6, 1.93, 2.4, 0.38, "nhiễu động khí quyển", size=13, color=INK, anchor="m")
    d2 = group(S)
    pic(d2.shapes, icon("wave", MUTED), 0.75, 3.75, 0.34, 0.34)
    text(d2.shapes, 1.13, 3.73, 1.9, 0.38, "rung động cơ", size=13, color=INK, anchor="m")
    d3 = group(S)
    pic(d3.shapes, icon("sync", MUTED), 4.65, 2.62, 0.34, 0.34)
    text(d3.shapes, 5.03, 2.5, 1.95, 0.6, "lượn vòng,\nđổi hướng nhanh", size=13, color=INK, anchor="m", line=1.0)
    good = line(S, tcx, tcy, tx, ty - 0.02, color=BLUE, w=3.0, tail="tri")
    good_lab = chip(S, 4.72, 4.2, 2.3, 0.44, "Gimbal giữ đường ngắm", fill=BLUE, size=12)
    # --- cot phai ---
    x0, w0 = 7.3, 5.5
    c1 = group(S)
    card(c1.shapes, x0, 1.38, w0, 1.62)
    label(c1.shapes, x0 + 0.2, 1.46, 5.0, "ĐẶC ĐIỂM UAV CÁNH BẰNG", color=BAD, ico="warning")
    text(c1.shapes, x0 + 0.22, 1.82, w0 - 0.4, 1.15, [
        {"t": "Bay nhanh, không treo tại chỗ, liên tục lượn vòng", "bullet": BAD},
        {"t": "Rung động cơ, nhiễu động khí quyển", "bullet": BAD},
        {"t": "Camera gắn cứng → ảnh nhòe, đường ngắm lệch", "bullet": BAD}],
        size=14.5, color=INK, space_after=3, line=1.0)
    c2 = group(S)
    card(c2.shapes, x0, 3.13, w0, 1.72)
    label(c2.shapes, x0 + 0.2, 3.2, 5.0, "ỨNG DỤNG HỆ ỔN ĐỊNH CAMERA", color=BLUE, ico="video")
    apps = [("binoculars", "Trinh sát,\ngiám sát"), ("crosshairs", "Bám, chỉ thị\nmục tiêu"),
            ("lifering", "Tuần tra,\ncứu nạn"), ("map", "Khảo sát,\nlập bản đồ")]
    tiles = []
    for i, (ic, lab) in enumerate(apps):
        tg = group(S)
        cx = x0 + 0.2 + i * 1.29
        box(tg.shapes, cx, 3.6, 1.2, 1.15, fill=TINT, radius=0.08)
        icon_disc(tg.shapes, cx + 0.6, 3.92, 0.46, icon(ic), BLUE)
        text(tg.shapes, cx + 0.02, 4.18, 1.16, 0.55, lab, size=12.5, color=INK, align="c", bold=True, line=0.95)
        tiles.append(tg)
    c3 = group(S)
    box(c3.shapes, x0, 4.98, w0, 1.64, fill=NAVY, radius=0.10, shd=True)
    text(c3.shapes, x0 + 0.25, 5.06, w0 - 0.5, 0.82,
         "→ Đề tài: **Thiết kế bộ điều khiển Gimbal trên cơ sở hệ điều hành thời gian thực NuttX**",
         size=16, color="FFFFFF", anchor="m", line=1.0)
    text(c3.shapes, x0 + 0.25, 5.86, w0 - 0.5, 0.72,
         "Apache NuttX: RTOS mã nguồn mở, chuẩn POSIX, là nền của PX4 trên UAV → tất định thời gian, dễ tích hợp",
         size=13, color="CFE3F7", anchor="m", line=1.0)
    a = Anim(s)
    a.auto().add(uav, "fly_l", 900)
    a.click().add(d1, "fade", 400).add(d2, "fade", 400, delay=150).add(d3, "fade", 400, delay=150)
    for b in bads:
        a.add(b, "wipe_l", 500, delay=120)
    a.add(bad_lab, "fade", 400, delay=200)
    a.add(c1, "rise", 500, delay=150)
    a.click().add(c2, "rise", 500)
    for t in tiles:
        a.add(t, "zoom", 350, delay=120)
    a.click()
    for b in bads:
        a.add(b, "exit_fade", 300)
    a.add(bad_lab, "exit_fade", 300)
    a.add(good, "wipe_l", 500, after=True).add(good_lab, "zoom", 400, delay=200)
    a.add(c3, "rise", 500, after=True).add(c3, "pulse", 300, after=True)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 2)


# =============================================================================
# 3. MUC TIEU VA YEU CAU
# =============================================================================
def s03():
    s = new_slide()
    S = s.shapes
    header(s, "MỞ ĐẦU · MỤC TIÊU VÀ YÊU CẦU", "Mục tiêu và yêu cầu của đồ án")
    obj = group(S)
    box(obj.shapes, LM, 1.38, CW, 1.22, fill=NAVY, radius=0.12, shd=True)
    icon_disc(obj.shapes, LM + 0.65, 1.99, 0.78, icon("bullseye"), ACCENT)
    text(obj.shapes, LM + 1.25, 1.42, CW - 1.45, 1.14,
         "**Mục tiêu:** thiết kế, chế tạo bộ điều khiển Gimbal hai trục Yaw – Pitch trên Apache NuttX, "
         "giữ ổn định đường ngắm camera khi khung mang chuyển động nhanh, đảo chiều và dừng đột ngột.",
         size=18, color="FFFFFF", anchor="m", line=1.02)
    lab = text(S, LM, 2.78, 9, 0.32, "YÊU CẦU · BỐN NHÓM CHỈ TIÊU ĐÁNH GIÁ", size=12.5, bold=True,
               color=BLUE, spc=40)
    reqs = [
        ("crosshairs", "Độ chính xác giữ hướng", ["σ khi giữ tĩnh", "*e*_{rms}, *e*_{max} khi khung chuyển động"]),
        ("chartline", "Đáp ứng động", ["Thời gian lên *t*_{r}, xác lập *t*_{s}", "Độ vọt lố, thời gian phục hồi"]),
        ("wave", "Khử nhiễu, bảo vệ cơ cấu", ["Tốc độ góc dư, tỉ lệ bão hòa *ρ*_{sat}", "Bước thay đổi lệnh, dòng động cơ"]),
        ("stopwatch", "Tất định thời gian", ["Vòng tốc độ góc **2 ms**", "Khe RS485 **5 ms** · mỗi trục **10 ms**"]),
    ]
    cards = []
    w = (CW - 3 * 0.2) / 4
    for i, (ic, ttl, items) in enumerate(reqs):
        g = group(S)
        x = LM + i * (w + 0.2)
        card(g.shapes, x, 3.12, w, 2.42)
        icon_disc(g.shapes, x + 0.48, 3.6, 0.66, icon(ic), BLUE if i < 3 else ACCENT)
        text(g.shapes, x + 0.92, 3.27, w - 1.0, 0.66, ttl, size=17, bold=True, color=NAVY, anchor="m", line=0.95)
        text(g.shapes, x + 0.22, 4.12, w - 0.34, 1.38, [{"t": t, "bullet": BLUE} for t in items],
             size=16, color=INK, space_after=6, line=1.0)
        cards.append(g)
    sc1 = chip(S, LM, 5.8, 6.0, 0.6, "Phạm vi: điều khiển Yaw và Pitch · Roll chỉ ước lượng", fill="FFFFFF",
               color=NAVY, size=15, line_=BLUE, radius=0.3)
    sc2 = chip(S, LM + 6.25, 5.8, 6.0, 0.6, "Phương pháp: lý thuyết → mô phỏng → thực nghiệm", fill="FFFFFF",
               color=NAVY, size=15, line_=BLUE, radius=0.3)
    a = Anim(s)
    a.click().add(obj, "wipe_l", 600)
    a.click().add(lab, "fade", 300)
    for c in cards:
        a.add(c, "rise", 450, delay=150)
    a.click().add(cards[3], "pulse", 300).add(sc1, "fade", 400, delay=200).add(sc2, "fade", 400, delay=150)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 3)


# =============================================================================
# 4. KET CAU DO AN
# =============================================================================
def s04():
    s = new_slide()
    S = s.shapes
    header(s, "MỞ ĐẦU", "Kết cấu đồ án tốt nghiệp")
    parts = [
        ("01", "Bài toán, yêu cầu và nguyên nhân sai lệch", "Gimbal Yaw–Pitch · chỉ tiêu · 5 nhóm nguyên nhân", "Chương 1"),
        ("02", "Thiết kế phần cứng và phần mềm thời gian thực", "Mạch · cơ khí · NuttX 7 luồng · 2 IMU", "Chương 2"),
        ("03", "Thuật toán điều khiển", "Cascade · bù IMU2 · bù trễ · giới hạn · đảo chiều · RS485", "Chương 2"),
        ("04", "Mô phỏng, hiệu chỉnh và thực nghiệm", "7 bước hiệu chỉnh, so sánh trước – sau trên hệ thật", "Chương 3"),
        ("05", "Đánh giá, kết luận và hướng phát triển", "Đạt · đánh đổi · hạn chế · hướng tiếp theo", "Kết luận"),
    ]
    ly = 2.55
    base = line(S, 1.7, ly, 11.65, ly, color="9FB8D3", w=3.0, tail=None)
    step = (11.65 - 1.7) / 4
    nodes = []
    for i, (num, ttl, sub, ch) in enumerate(parts):
        cx = 1.7 + i * step
        hot = num == "04"
        d = 1.0 if hot else 0.86
        g = group(S)
        box(g.shapes, cx - d / 2, ly - d / 2, d, d, fill=ACCENT if hot else NAVY, kind=MSO_SHAPE.OVAL,
            line="FFFFFF", lw=3, text=num, size=24 if hot else 21, bold=True, color="FFFFFF")
        card(g.shapes, cx - 1.14, 3.35, 2.28, 2.85, fill="FFFFFF", line_=ACCENT if hot else None, lw=2.0)
        text(g.shapes, cx - 1.04, 3.48, 2.08, 1.22, ttl, size=16, bold=True, color=NAVY, align="c", anchor="m",
             line=0.98)
        text(g.shapes, cx - 1.04, 4.72, 2.08, 0.9, sub, size=13, color=MUTED, align="c", line=0.98)
        chip(g.shapes, cx - 0.7, 5.68, 1.4, 0.36, ch, fill=ACCENT if hot else BLUE, size=12.5)
        if hot:
            chip(g.shapes, cx - 0.72, 1.62, 1.44, 0.36, "TRỌNG TÂM", fill=ACCENT, size=12.5)
        nodes.append(g)
    a = Anim(s)
    a.auto().add(base, "wipe_l", 700)
    for n in nodes:
        a.add(n, "rise", 450, after=True)
    a.add(nodes[3], "pulse", 300, after=True)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 4)


# =============================================================================
# 5. BAI TOAN ON DINH GIMBAL HAI TRUC
# =============================================================================
def s05():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 1 · BÀI TOÁN VÀ YÊU CẦU", "Bài toán: Gimbal tạo chuyển động bù để đường ngắm đứng yên")
    card(S, LM, 1.38, 6.3, 5.24)
    iw = 6.0
    ip = pic(S, img("gimbal_illus"), 0.7, 1.5, iw)
    sc = iw / 1073.0
    rings = []
    for (px, py) in ((538, 152 - 70), (662, 560 - 70)):
        cx, cy = 0.7 + px * sc, 1.5 + py * sc
        rings.append(box(S, cx - 0.33, cy - 0.33, 0.66, 0.66, kind=MSO_SHAPE.OVAL, line=ACCENT, lw=3.0))
    chain = [("Khung mang", "7F8B98"), ("Trục Yaw · ~~*q*_{y}~~", YAW), ("Trục Pitch · ~~*q*_{p}~~", PITCH),
             ("Camera", NAVY)]
    ws = [1.38, 1.45, 1.55, 1.0]
    x = 0.72
    chain_sh = []
    for i, ((t, col), w) in enumerate(zip(chain, ws)):
        g = group(S)
        box(g.shapes, x, 5.32, w, 0.5, fill=col, radius=0.1, text=t, size=13.5, bold=True, color="FFFFFF",
            margin=0.03)
        if i < 3:
            arrow(g.shapes, x + w + 0.03, 5.57, x + w + 0.18, 5.57, color=MUTED, w=2.0)
        chain_sh.append(g)
        x += w + 0.21
    clab = text(S, 0.72, 5.92, 6.0, 0.3, "Chuỗi động học nối tiếp: động cơ Pitch nằm trên khung Yaw",
                size=12.5, color=MUTED, italic=True)
    X, W = 7.05, 5.75
    pr = group(S)
    card(pr.shapes, X, 1.38, W, 1.58)
    label(pr.shapes, X + 0.2, 1.46, W - 0.4, "NGUYÊN LÝ ỔN ĐỊNH", ico="bullseye")
    text(pr.shapes, X + 0.22, 1.82, W - 0.44, 1.1,
         "Góc đường ngắm = góc khung mang + góc quay các khớp ⇒ Gimbal quay bù **ngược** chuyển động "
         "khung mang để đường ngắm đứng yên trong không gian", size=15, color=INK, line=1.02)
    imu = []
    for i, (ic, col, ttl, body) in enumerate([
            ("video", BLUE, "IMU1 · trên camera", "Đo ~~*θ*~~, ~~*ω*~~ của đường ngắm → tín hiệu **phản hồi**"),
            ("compass", NAVY, "IMU2 · trên khung mang", "Đo ~~*ω*_{b}~~, ~~*α*_{b}~~ của khung → **bù trước**")]):
        g = group(S)
        x = X + i * (W / 2 + 0.08)
        w = W / 2 - 0.08
        card(g.shapes, x, 3.1, w, 1.55)
        icon_disc(g.shapes, x + 0.42, 3.5, 0.56, icon(ic), col)
        text(g.shapes, x + 0.78, 3.22, w - 0.85, 0.56, ttl, size=15, bold=True, color=NAVY, anchor="m")
        text(g.shapes, x + 0.2, 3.86, w - 0.35, 0.75, body, size=14, color=INK, line=1.0)
        imu.append(g)
    ch = group(S)
    card(ch.shapes, X, 4.8, W, 1.82)
    label(ch.shapes, X + 0.2, 4.88, W - 0.4, "TÌNH HUỐNG KHÓ", color=BAD, ico="warning")
    text(ch.shapes, X + 0.22, 5.24, W - 0.44, 1.35, [
        {"t": "Khung quay nhanh (giá thử đến ≈ 240 °/s), đảo chiều, dừng đột ngột", "bullet": BAD},
        {"t": "Trễ ≈ 12 ms trên đường tín hiệu; bus RS485 dùng chung", "bullet": BAD},
        {"t": "Yêu cầu: sai lệch nhỏ trên **cả Yaw và Pitch**", "bullet": BAD}],
        size=14.5, color=INK, space_after=2, line=1.0)
    a = Anim(s)
    a.click()
    for i, g in enumerate(chain_sh):
        a.add(g, "wipe_l", 350, after=i > 0)
    a.add(clab, "fade", 300, after=True)
    a.click().add(pr, "rise", 500)
    a.click().add(imu[0], "rise", 450).add(rings[0], "zoom", 400, delay=150)
    a.add(imu[1], "rise", 450, delay=250).add(rings[1], "zoom", 400, delay=150)
    a.click().add(ch, "rise", 500)
    a.build()
    set_transition(s, "push", dir="u")
    number(s)
    notes(s, 5)


# =============================================================================
# 6. NGUYEN NHAN SAI LECH
# =============================================================================
def s06():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 1 · BÀI TOÁN VÀ YÊU CẦU", "Nguyên nhân sai lệch: phản hồi đến muộn, trễ chặn việc tăng hệ số")
    # A - sai lech khi chi phan hoi
    A = group(S)
    card(A.shapes, LM, 1.38, 3.95, 2.62)
    label(A.shapes, LM + 0.18, 1.45, 3.7, "CHỈ CÓ PHẢN HỒI (1.28)", ico="warning", color=BAD)
    text(A.shapes, LM + 0.1, 1.86, 3.75, 0.7, "~~*e*_{θ,ss} = −*ω*_{b} / *K*_{p}^{θ}~~", size=27, color=NAVY,
         align="c", anchor="m")
    text(A.shapes, LM + 0.2, 2.62, 3.6, 0.62,
         "Khung quay 8 °/s, ~~*K*_{p}^{θ}~~ = 9,5 s^{−1} → sai lệch ≈ **0,84°**", size=14.5, color=INK, line=1.0)
    chip(A.shapes, LM + 0.2, 3.3, 3.55, 0.5, "Giảm 10 lần ⇒ phải tăng ~~*K*_{p}^{θ}~~ 10 lần", fill=ACC_TINT,
         color=INK, size=14, radius=0.1)
    # B - nguon tre
    B = group(S)
    bx, bw = 4.7, 4.6
    card(B.shapes, bx, 1.38, bw, 2.62)
    label(B.shapes, bx + 0.18, 1.45, bw - 0.3, "CÁC NGUỒN TRỄ (BẢNG 1.4)", ico="stopwatch")
    x0, sc = 6.62, 0.125
    for t in (0, 5, 10, 15):
        text(B.shapes, x0 + t * sc - 0.3, 1.8, 0.6, 0.22, f"{t}" + (" ms" if t == 15 else ""), size=11,
             color=MUTED, align="c")
        line(B.shapes, x0 + t * sc, 2.02, x0 + t * sc, 3.48, color="E3E8EE", w=0.75, tail=None)
    bars = []
    rows = [("~~*τ*_{s}~~ lấy mẫu, lọc IMU", 2, 4, "2–4 ms"), ("~~*τ*_{c}~~ tính toán", 0, 2, "< 2 ms"),
            ("~~*τ*_{b}~~ chờ khe RS485", 5, 10, "5–10 ms"), ("~~*τ*_{m}~~ đáp ứng động cơ", 8, 14, "8–14 ms")]
    for i, (nm, a0, a1, vtxt) in enumerate(rows):
        y = 2.07 + i * 0.36
        g = group(S)
        text(g.shapes, bx + 0.18, y - 0.03, 1.95, 0.34, nm, size=12.5, color=INK, anchor="m")
        box(g.shapes, x0 + a0 * sc, y + 0.04, (a1 - a0) * sc, 0.24, fill=ACCENT if i >= 2 else BLUE, radius=0.06)
        text(g.shapes, x0 + a1 * sc + 0.05, y - 0.03, 0.9, 0.34, vtxt, size=12, color=INK, anchor="m", bold=True)
        bars.append(g)
    tsum = text(S, bx + 0.18, 3.52, bw - 0.3, 0.42, "→ trễ hiệu dụng cần bù ~~*τ*_{p}~~ ≈ **12 ms**", size=15,
                color=NAVY, anchor="m")
    # C - mat pha
    Cg = group(S)
    cx, cw = 9.5, 3.3
    card(Cg.shapes, cx, 1.38, cw, 2.62)
    label(Cg.shapes, cx + 0.18, 1.45, cw - 0.3, "TRỄ LÀM MẤT PHA (1.31)", ico="wave")
    text(Cg.shapes, cx + 0.1, 1.84, cw - 0.2, 0.5, "~~Δ*φ* = −360°·*f*·*τ*_{p}~~", size=19, color=NAVY, align="c",
         anchor="m")
    for i, (v, f) in enumerate((("−43°", "tại 10 Hz"), ("−86°", "tại 20 Hz"))):
        text(Cg.shapes, cx + 0.15 + i * 1.5, 2.4, 1.5, 0.5, v, size=26, bold=True, color=ACCENT, align="c", anchor="m")
        text(Cg.shapes, cx + 0.15 + i * 1.5, 2.9, 1.5, 0.3, f, size=12.5, color=INK, align="c")
    text(Cg.shapes, cx + 0.15, 3.32, cw - 0.3, 0.55, "→ không thể tăng hệ số tùy ý", size=14, bold=True, color=BAD,
         align="c", anchor="m")
    # Hang duoi - 5 nguyen nhan
    lab = text(S, LM, 4.13, 11, 0.32, "5 NHÓM NGUYÊN NHÂN → ẢNH HƯỞNG → HƯỚNG GIẢI PHÁP", size=12.5, bold=True,
               color=BLUE, spc=40)
    causes = [
        ("wave", "Nhiễu đo IMU", "Rung lệnh; nhiễu bị khuếch đại khi tăng hệ số", "Lọc IMU · ① vòng trong nhanh"),
        ("stopwatch", "Trễ và biên ổn định", "Lệnh bù đến muộn, không tăng được hệ số", "② bù IMU2 · ③ ngoại suy"),
        ("network", "Truyền thông không tất định", "Trễ ngẫu nhiên, cập nhật các trục không đều", "⑥ lịch RS485 luân phiên"),
        ("route", "Chuyển động nhanh, đảo chiều", "Bão hòa, rung; camera vượt ngược khi dừng", "⑤ đảo chiều, dừng, tạo dạng"),
        ("sliders", "Giới hạn lệnh cố định", "Thấp: bão hòa · cao: lệnh lớn khi tĩnh", "④ giới hạn theo trạng thái"),
    ]
    w = (CW - 4 * 0.15) / 5
    tiles, sols = [], []
    for i, (ic, ttl, eff, sol) in enumerate(causes):
        x = LM + i * (w + 0.15)
        g = group(S)
        card(g.shapes, x, 4.5, w, 2.12, fill="FFFFFF")
        icon_disc(g.shapes, x + 0.36, 4.86, 0.48, icon(ic), BAD)
        text(g.shapes, x + 0.64, 4.56, w - 0.68, 0.62, ttl, size=13, bold=True, color=NAVY, anchor="m", line=0.95)
        text(g.shapes, x + 0.15, 5.2, w - 0.27, 0.7, eff, size=13, color=INK, line=0.98)
        tiles.append(g)
        sols.append(chip(S, x + 0.1, 5.97, w - 0.2, 0.55, sol, fill=NAVY, size=12.5, radius=0.1))
    a = Anim(s)
    a.click().add(A, "rise", 500)
    a.click().add(B, "fade", 400)
    for b in bars:
        a.add(b, "wipe_l", 400, after=True)
    a.add(tsum, "fade", 400, after=True)
    a.click().add(Cg, "rise", 500)
    a.click().add(lab, "fade", 300)
    for t in tiles:
        a.add(t, "rise", 400, delay=150)
    a.click()
    for c in sols:
        a.add(c, "zoom", 350, delay=150)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 6)


# =============================================================================
# 7. KIEN TRUC PHAN CUNG
# =============================================================================
def s07():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 2 · THIẾT KẾ PHẦN CỨNG", "Phần cứng: 2 IMU trên 2 bus SPI, 2 động cơ chung bus RS485")
    D1 = group(S)
    card(D1.shapes, LM, 1.38, 3.3, 3.62)
    label(D1.shapes, LM + 0.15, 1.45, 3.0, "MIỀN CẢM BIẾN", ico="compass")
    pic(D1.shapes, img("icm20948"), 1.42, 1.82, 1.55)
    for i, (t1, t2) in enumerate((("**IMU1** · trên camera", "ICM20948 · SPI1 · 500 Hz"),
                                  ("**IMU2** · trên khung mang", "ICM20948 · SPI2 · 500 Hz"))):
        box(D1.shapes, 0.7, 2.98 + i * 0.97, 3.0, 0.82, fill=TINT, radius=0.08,
            text=[{"t": t1, "size": 14.5}, {"t": t2, "size": 12.5, "color": MUTED}], color=INK)
    D2 = group(S)
    card(D2.shapes, 4.45, 1.38, 3.9, 3.62)
    label(D2.shapes, 4.6, 1.45, 3.6, "MIỀN XỬ LÝ", ico="chip")
    pic(D2.shapes, img("stm32"), 5.85, 1.8, 1.12)
    box(D2.shapes, 4.6, 2.98, 3.6, 0.82, fill=TINT, radius=0.08,
        text=[{"t": "**STM32F407VET6**", "size": 15}, {"t": "Cortex-M4F · 168 MHz · FPU", "size": 12.5,
                                                         "color": MUTED}], color=INK)
    box(D2.shapes, 4.6, 3.95, 3.6, 0.82, fill=NAVY, radius=0.08,
        text=[{"t": "**Apache NuttX**", "size": 15}, {"t": "7 luồng thời gian thực · SCHED_FIFO", "size": 12.5,
                                                       "color": "CFE3F7"}], color="FFFFFF")
    D3 = group(S)
    card(D3.shapes, 8.95, 1.38, 3.85, 3.62)
    label(D3.shapes, 9.1, 1.45, 3.6, "MIỀN TRUYỀN ĐỘNG", ico="cog")
    pic(D3.shapes, img("ms3506"), 9.9, 1.8, 1.95)
    box(D3.shapes, 9.1, 2.98, 3.55, 0.6, fill=TINT, radius=0.08, text="**MAX485** · RS485 bán song công",
        size=13, color=INK)
    line(D3.shapes, 9.55, 3.77, 12.2, 3.77, color=BLUE, w=3.0, tail=None)
    line(D3.shapes, 10.88, 3.58, 10.88, 3.77, color=BLUE, w=2.0, tail=None)
    for i, (nm, col, idn) in enumerate((("Yaw", YAW, "ID 1"), ("Pitch", PITCH, "ID 2"))):
        x = 9.1 + i * 1.85
        line(D3.shapes, x + 0.85, 3.77, x + 0.85, 3.95, color=BLUE, w=2.0, tail=None)
        box(D3.shapes, x, 3.95, 1.7, 0.82, fill=col, radius=0.08,
            text=[{"t": f"**MS3506 · {nm}**", "size": 13.5}, {"t": idn, "size": 12.5}], color="FFFFFF")
    ar1 = group(S)
    for i, t in enumerate(("SPI1", "SPI2")):
        y = 3.39 + i * 0.97
        arrow(ar1.shapes, 3.72, y, 4.58, y, color=BLUE, w=2.25)
        text(ar1.shapes, 3.72, y - 0.33, 0.86, 0.28, t, size=12, bold=True, color=BLUE, align="c")
    ar2 = group(S)
    arrow(ar2.shapes, 8.22, 3.28, 9.08, 3.28, color=BLUE, w=2.25)
    text(ar2.shapes, 8.18, 2.95, 0.95, 0.28, "UART5", size=12, bold=True, color=BLUE, align="c")
    P = group(S)
    card(P.shapes, LM, 5.18, 5.55, 1.1)
    label(P.shapes, LM + 0.15, 5.24, 2.5, "NGUỒN", ico="bolt", color=ACCENT)
    x = 0.72
    for i, (t, w) in enumerate((("12 V", 0.62), ("MP1584EN", 1.2), ("5 V", 0.55), ("LM1117", 0.95), ("3,3 V", 0.68))):
        stage = i % 2 == 1
        chip(P.shapes, x, 5.66, w, 0.42, t, fill="FFFFFF" if stage else NAVY, color=NAVY if stage else "FFFFFF",
             line_=NAVY if stage else None, size=12.5, radius=0.08)
        if i < 4:
            arrow(P.shapes, x + w + 0.03, 5.87, x + w + 0.2, 5.87, color=MUTED, w=1.75)
        x += w + 0.24
    M = group(S)
    card(M.shapes, 6.3, 5.18, 6.5, 1.1)
    label(M.shapes, 6.45, 5.24, 3.0, "KÊNH GIÁM SÁT", ico="desktop")
    chip(M.shapes, 6.48, 5.66, 0.95, 0.42, "USART", fill=NAVY, size=12.5, radius=0.08)
    arrow(M.shapes, 7.46, 5.87, 7.66, 5.87, color=MUTED)
    pic(M.shapes, img("usbttl"), 7.72, 5.6, 0.62, 0.62)
    text(M.shapes, 8.38, 5.66, 1.05, 0.42, "USB–TTL", size=12.5, bold=True, color=NAVY, anchor="m")
    arrow(M.shapes, 9.45, 5.87, 9.65, 5.87, color=MUTED)
    chip(M.shapes, 9.7, 5.66, 2.95, 0.42, "Máy tính · phần mềm giám sát", fill=BLUE, size=12.5, radius=0.08)
    pr = text(S, LM, 6.34, CW, 0.3, "Nguyên tắc: tách hai đường cảm biến (SPI1 / SPI2) · tách kênh điều khiển và kênh "
              "giám sát · chiều truyền bus RS485 do driver điều khiển", size=12.5, color=MUTED, italic=True)
    a = Anim(s)
    a.click().add(D1, "rise", 500)
    a.click().add(ar1, "wipe_l", 400).add(D2, "rise", 500, after=True)
    a.click().add(ar2, "wipe_l", 400).add(D3, "rise", 500, after=True)
    a.click().add(P, "rise", 450).add(M, "rise", 450, delay=200).add(pr, "fade", 400, after=True)
    a.build()
    set_transition(s, "push", dir="u")
    number(s)
    notes(s, 7)


# =============================================================================
# 8. THIET KE MACH VA KHUNG CO KHI
# =============================================================================
def img_card(S, x, y, w, h, path, cap, pad=0.1):
    g = group(S)
    card(g.shapes, x, y, w, h)
    iw, ih = Image.open(path).size
    aw, ah = w - 2 * pad, h - 2 * pad - 0.36
    sc = min(aw / iw, ah / ih)
    pw, ph = iw * sc, ih * sc
    pic(g.shapes, path, x + (w - pw) / 2, y + pad + (ah - ph) / 2, pw, ph)
    text(g.shapes, x + 0.05, y + h - 0.4, w - 0.1, 0.34, cap, size=12.5, color=INK, align="c", anchor="m")
    return g


def s08():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 2 · THIẾT KẾ PHẦN CỨNG", "Thiết kế mạch điều khiển 5 khối và khung cơ khí hai trục")
    rows = []
    for r, (ic, t1, t2, y, h, items) in enumerate([
        ("chip", "Mạch điều khiển", "Altium · 2 lớp", 1.38, 2.32,
         [("sch_mcu", "Sơ đồ nguyên lý"), ("pcb_layout", "Mạch in 2 lớp"), ("pcb_3d", "Mô hình 3D"),
          ("pcb_real", "Mạch đã chế tạo")]),
        ("cube", "Khung cơ khí", "12 chi tiết", 4.32, 2.3,
         [("cad_2d", "Bản vẽ hai hình chiếu"), ("cad_exploded", "Kết cấu tách rời"),
          ("cad_render", "Mô hình lắp ráp"), (None, None)]),
    ]):
        lab = group(S)
        icon_disc(lab.shapes, 1.3, y + 0.55, 0.66, icon(ic), BLUE if r == 0 else NAVY)
        text(lab.shapes, LM, y + 0.95, 1.6, 0.62, t1, size=16, bold=True, color=NAVY, align="c", line=0.95)
        text(lab.shapes, LM, y + 1.55, 1.6, 0.3, t2, size=13, color=MUTED, align="c")
        items_sh = [lab]
        for i, (nm, cap) in enumerate(items):
            x = 2.3 + i * 2.7
            if nm is None:
                g = group(S)
                card(g.shapes, x, y, 2.4, h, fill=TINT)
                text(g.shapes, x + 0.12, y + 0.1, 2.2, h - 0.2, [
                    {"t": "**Kết cấu chính**", "size": 13.5, "color": NAVY},
                    {"t": "Đế carbon, 4 đệm cao su; IMU2 trên đế", "bullet": BLUE},
                    {"t": "Trục Yaw: MS3506 + tấm quay", "bullet": BLUE},
                    {"t": "Trục Pitch: MS3506 + mặt bích, giá chữ L mang camera, IMU1", "bullet": BLUE}],
                    size=12.5, color=INK, space_after=2, line=0.95, anchor="m")
                items_sh.append(g)
                continue
            items_sh.append(img_card(S, x, y, 2.4, h, img(nm), cap))
            if i < 3:
                arrow(S, x + 2.43, y + h / 2, x + 2.67, y + h / 2, color=ACCENT, w=2.25)
        rows.append(items_sh)
    chips = group(S)
    x = 2.3
    for t, w in (("STM32F407VET6", 1.85), ("2 cổng SPI cảm biến", 2.1), ("RS485 · MAX485", 1.65),
                 ("Chuyển mức 4×TXS0102", 2.3), ("Nguồn 12 → 5 → 3,3 V", 2.1)):
        chip(chips.shapes, x, 3.8, w, 0.38, t, fill="FFFFFF", color=NAVY, line_=BLUE, size=12)
        x += w + 0.125
    a = Anim(s)
    a.click()
    for i, g in enumerate(rows[0]):
        a.add(g, "wipe_l" if i else "fade", 450, after=i > 0)
    a.add(chips, "fade", 400, after=True)
    a.click()
    for i, g in enumerate(rows[1]):
        a.add(g, "wipe_l" if i else "fade", 450, after=i > 0)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 8)


# =============================================================================
# 9. SAN PHAM DA CHE TAO
# =============================================================================
def s09():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 2 · SẢN PHẨM ĐÃ CHẾ TẠO", "Sản phẩm: Gimbal, mạch điều khiển, 2 × MS3506, 2 × ICM20948")
    hero = group(S)
    card(hero.shapes, LM, 1.38, 5.75, 5.24)
    ph = 4.72
    pw = ph * 1126 / 1230
    px0, py0 = LM + (5.75 - pw) / 2, 1.5
    pic(hero.shapes, img("gimbal_photo"), px0, py0, pw, ph)
    text(hero.shapes, LM + 0.1, 6.26, 5.55, 0.32, "Cơ cấu Gimbal hai trục hoàn chỉnh (Hình 3.1)", size=12.5,
         color=MUTED, align="c", italic=True)
    sc = pw / 563.0
    calls = [  # (diem tren anh goc 563x615), (vi tri nhan), noi dung
        ((452, 135), (3.65, 1.6, 2.3), "Động cơ Pitch · MS3506"),
        ((178, 238), (0.72, 2.3, 1.9), "Giá camera + IMU1"),
        ((262, 350), (4.0, 4.18, 2.2), "Động cơ Yaw · MS3506"),
        ((232, 478), (3.62, 5.66, 2.4), "IMU2 · ICM20948 trên đế"),
        ((112, 452), (0.68, 5.66, 2.15), "Đế carbon, đệm cao su"),
    ]
    call_sh = []
    for (qx, qy), (lx, ly, lw), t in calls:
        g = group(S)
        tx, ty = px0 + qx * sc, py0 + qy * sc
        ex = lx + lw / 2 if abs((lx + lw / 2) - tx) < 1.0 else (lx if lx > tx else lx + lw)
        ey = ly + 0.17
        line(g.shapes, ex, ey, tx, ty, color=ACCENT, w=1.75, tail="oval", size="sm")
        chip(g.shapes, lx, ly, lw, 0.36, t, fill=NAVY, size=12, radius=0.08)
        call_sh.append(g)
    X, W = 6.5, 6.3
    cards = []
    for i, (nm, ttl, l1, l2, iw) in enumerate([
            ("pcb_real", "Mạch điều khiển", "STM32F407VET6 · 2 cổng SPI · RS485 · nguồn 12 V",
             "Thiết kế trên Altium, mạch in 2 lớp", 1.3),
            ("ms3506", "2 × động cơ MS3506", "BLDC tích hợp driver, giao tiếp RS485",
             "ID 1: trục Yaw · ID 2: trục Pitch", 2.0),
            ("icm20948", "2 × cảm biến ICM20948", "IMU 9 trục, đọc qua SPI, tần số 500 Hz",
             "IMU1: giá camera · IMU2: đế khung mang", 1.85)]):
        y = 1.38 + i * 1.77
        g = group(S)
        card(g.shapes, X, y, W, 1.62)
        path = img(nm)
        iw0, ih0 = Image.open(path).size
        ih = min(1.42, iw * ih0 / iw0)
        iw2 = ih * iw0 / ih0
        pic(g.shapes, path, X + 0.1 + (2.1 - iw2) / 2, y + (1.62 - ih) / 2, iw2, ih)
        text(g.shapes, X + 2.35, y + 0.12, W - 2.5, 0.45, ttl, size=17, bold=True, color=NAVY, anchor="m")
        text(g.shapes, X + 2.35, y + 0.6, W - 2.5, 0.95, [{"t": l1, "bullet": BLUE}, {"t": l2, "bullet": BLUE}],
             size=13.5, color=INK, space_after=2, line=1.0)
        cards.append(g)
    a = Anim(s)
    a.click().add(hero, "zoom", 600)
    for g in call_sh:
        a.add(g, "fade", 350, after=True)
    a.click()
    for i, g in enumerate(cards):
        a.add(g, "rise", 450, delay=0 if i == 0 else 200)
    a.build()
    set_transition(s, "zoom")
    number(s)
    notes(s, 9)


# =============================================================================
# 10. PHAN MEM NUTTX
# =============================================================================
def s10():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 2 · TỔ CHỨC PHẦN MỀM TRÊN APACHE NUTTX",
           "Apache NuttX: 7 luồng ưu tiên chiếm quyền giữ nhịp 2 ms")
    text(S, LM, 1.38, 4.2, 0.32, "TỔ CHỨC 5 TẦNG", size=12.5, bold=True, color=BLUE, spc=40)
    layers = [
        ("Ứng dụng", ACCENT, "7 pthread của gimbal + luồng console"),
        ("Đồng bộ", "3D7CC4", "semaphore imu_ready, safety · mutex mẫu IMU, hộp thư lệnh, bus RS485"),
        ("Nhân NuttX", BLUE, "FLAT build · SCHED_FIFO 1–255 · mutex thừa kế ưu tiên"),
        ("Driver /dev", "123A6B", "/dev/spi1, spi2, ttyS1, gpio0–2, console"),
        ("Phần cứng", NAVY, "STM32F407: SPI1, SPI2, UART5, GPIO, USART"),
    ]
    lay_sh = []
    for i, (nm, col, desc) in enumerate(layers):
        y = 1.76 + i * 0.97
        g = group(S)
        box(g.shapes, LM, y, 1.3, 0.88, fill=col, radius=0.08, text=nm, size=13.5, bold=True, color="FFFFFF")
        card(g.shapes, LM + 1.38, y, 2.8, 0.88, shd=False, line_=LINE_SOFT)
        text(g.shapes, LM + 1.48, y, 2.62, 0.88, desc, size=12.5, color=INK, anchor="m", line=0.95)
        lay_sh.append(g)
    # gian do thoi gian
    X, W = 4.95, 7.85
    card(S, X, 1.38, W, 4.2)
    text(S, X + 0.18, 1.44, W - 0.3, 0.32, "GIẢN ĐỒ THỜI GIAN 0 – 10 ms · LẬP LỊCH ƯU TIÊN CHIẾM QUYỀN",
         size=12.5, bold=True, color=BLUE, spc=30)
    x0, k = 7.35, 0.52
    axis = group(S)
    for t in range(0, 11, 2):
        text(axis.shapes, x0 + t * k - 0.3, 1.8, 0.6, 0.25, f"{t} ms" if t in (0, 10) else str(t), size=11.5,
             color=MUTED, align="c")
        line(axis.shapes, x0 + t * k, 2.07, x0 + t * k, 5.36, color="E3E8EE", w=0.75, tail=None)
    names = ["IMU · 120", "Giám sát an toàn · 118", "Điều khiển · 115", "RS485 · 110→116→114",
             "Khôi phục IMU2 · 108", "Báo cáo · 90", "LED · 50"]
    ry = [2.1 + i * 0.465 for i in range(7)]
    for i, nm in enumerate(names):
        text(axis.shapes, X + 0.18, ry[i], 2.25, 0.42, nm, size=12.5, color=INK, anchor="m",
             bold=i in (0, 2, 3))
    imu_bars = group(S)
    ctl_bars = group(S)
    for t in range(0, 10, 2):
        box(imu_bars.shapes, x0 + t * k, ry[0] + 0.09, 0.35 * k, 0.26, fill=BLUE, radius=0.04)
        box(ctl_bars.shapes, x0 + (t + 0.35) * k, ry[2] + 0.09, 0.65 * k, 0.26, fill=ACCENT, radius=0.04)
    text(ctl_bars.shapes, x0 + 1.05 * k, ry[2] + 0.02, 0.95 * k, 0.4, "θ+ω", size=11, color=INK, anchor="m")
    sem = group(S)
    arrow(sem.shapes, x0 + 0.17 * k, ry[0] + 0.36, x0 + 0.45 * k, ry[2] + 0.08, color=NAVY, w=1.5)
    text(sem.shapes, x0 + 0.55 * k, ry[1] + 0.02, 2.4, 0.4, "semaphore · hạn chờ 6 ms", size=11.5, color=NAVY,
         italic=True, anchor="m")
    rs = group(S)
    for j, (nm, col) in enumerate((("Yaw", YAW), ("Pitch", PITCH))):
        xs = x0 + j * 5 * k
        box(rs.shapes, xs + 0.02, ry[3] + 0.06, 5 * k - 0.04, 0.32, fill="FFFFFF", line=col, lw=1.25, radius=0.05)
        box(rs.shapes, xs + 0.02, ry[3] + 0.06, 2.25 * k, 0.32, fill=col, radius=0.05, text=nm, size=11.5,
            bold=True, color="FFFFFF")
        text(rs.shapes, xs + 2.3 * k, ry[3] + 0.03, 2.6 * k, 0.38, "khe 5 ms", size=11, color=MUTED, anchor="m")
    other = group(S)
    text(other.shapes, x0 + 0.1, ry[4], 4.9, 0.42, "chu kỳ 50 ms (nền)", size=11.5, color=MUTED, italic=True,
         anchor="m")
    box(other.shapes, x0 + 1.0 * k, ry[5] + 0.09, 0.4 * k, 0.26, fill="9AA7B4", radius=0.04)
    text(other.shapes, x0 + 1.5 * k, ry[5], 3.0, 0.42, "chu kỳ 10 ms", size=11.5, color=MUTED, italic=True,
         anchor="m")
    text(other.shapes, x0 + 0.1, ry[6], 4.9, 0.42, "chu kỳ 50 ms", size=11.5, color=MUTED, italic=True, anchor="m")
    text(other.shapes, x0 + 2.95, ry[1], 2.25, 0.42, "theo sự kiện (lỗi, mất mẫu)", size=11.5, color=MUTED,
         italic=True, anchor="m")
    ph = line(S, x0, 2.02, x0, 5.38, color=BAD, w=2.25, tail=None)
    dec = []
    for i, t in enumerate(("Chờ semaphore có hạn 6 ms → quá hạn thì kiểm tra tuổi mẫu",
                           "Chỉ dùng mẫu mới nhất → trễ không tích lũy",
                           "RS485 nâng ưu tiên lên 116 khi phát khung")):
        w = (W - 0.3) / 3
        g = group(S)
        card(g.shapes, X + i * (w + 0.15), 5.72, w, 0.9, fill=ACC_TINT, shd=False)
        pic(g.shapes, icon("check", ACCENT), X + i * (w + 0.15) + 0.12, 5.98, 0.36, 0.36)
        text(g.shapes, X + i * (w + 0.15) + 0.55, 5.72, w - 0.62, 0.9, t, size=13, color=INK, anchor="m", line=0.98)
        dec.append(g)
    a = Anim(s)
    a.click()
    for i, g in enumerate(reversed(lay_sh)):
        a.add(g, "rise", 350, delay=0 if i == 0 else 130)
    a.click().add(imu_bars, "wipe_l", 700).add(sem, "fade", 400, delay=250)
    a.add(ctl_bars, "wipe_l", 700, delay=150)
    a.add(ph, "appear", 1, after=True).add(ph, "path", 4000, path=f"M 0 0 L {10 * k / 13.333:.4f} 0 E", ease=False)
    a.click().add(rs, "wipe_l", 700).add(other, "fade", 500, delay=200)
    a.click()
    for i, g in enumerate(dec):
        a.add(g, "rise", 400, delay=0 if i == 0 else 150)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 10)


# =============================================================================
# 11. HAI IMU - XU LY, LOC, UOC LUONG TU THE
# =============================================================================
def s11():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 2 · HỆ THỐNG HAI CẢM BIẾN IMU",
           "Hai IMU: IMU1 cho phản hồi, IMU2 đo khung mang để bù trước")
    text(S, LM, 1.36, 6, 0.3, "CHUỖI XỬ LÝ TRONG MỖI CHU KỲ 2 ms", size=12.5, bold=True, color=BLUE, spc=40)
    steps = [("Đọc 2 IMU", "SPI1, SPI2 · 500 Hz"), ("Kiểm tra hợp lệ", "tuổi mẫu ≤ 6 ms"),
             ("Đồng bộ 2 mẫu", "lệch ≤ 1,5 ms"), ("Hiệu chuẩn", "tỉ lệ, lệch trục, điểm 0"),
             ("Lọc", "con quay 80 Hz · ~~*α*_{b}~~ 18 Hz"), ("Mahony + Kalman", "→ ~~*θ*~~, ~~*ω*~~ · phát semaphore")]
    cols = ["123A6B", "1A4F8B", BLUE, "2F78C0", "3D7CC4", NAVY]
    chev = []
    bw, gap = 1.88, 0.194
    for i, ((t1, t2), col) in enumerate(zip(steps, cols)):
        x = LM + i * (bw + gap)
        g = group(S)
        box(g.shapes, x, 1.72, bw, 1.08, fill=col, radius=0.1,
            text=[{"t": t1, "size": 14, "bold": True}, {"t": t2, "size": 12}], color="FFFFFF", margin=0.07,
            line=0.95)
        if i < len(steps) - 1:
            arrow(g.shapes, x + bw + 0.01, 2.26, x + bw + gap - 0.01, 2.26, color=ACCENT, w=2.25, size="sm")
        chev.append(g)
    cards = []
    for i, (x, w, ic, col, ttl, items, foot, ftint, fcol) in enumerate([
        (LM, 5.25, "video", BLUE, "IMU1 · trên giá camera", [
            "Mahony + Kalman → góc ~~*θ*~~, tốc độ góc ~~*ω*~~ của đường ngắm",
            "Hằng số thời gian tư thế từ gia tốc kế 0,45 s",
            "Tín hiệu phản hồi cho cả hai vòng"], "Mất IMU1 → vòng kín mất phản hồi → dừng an toàn", BAD_TINT, BAD),
        (6.0, 6.8, "compass", NAVY, "IMU2 · trên đế (khung mang)", [
            "Tốc độ góc khung ~~*ω*_{b}~~ lọc 60 Hz (Yaw), 55 Hz (Pitch)",
            "Gia tốc góc ~~*α*_{b}~~ từ vi phân, lọc 18 Hz",
            "Chiếu lên trục Pitch theo góc Yaw (2.9):"], "Mất IMU2 → chỉ mất khâu bù, vòng kín vẫn làm việc",
         TINT, NAVY)]):
        g = group(S)
        card(g.shapes, x, 3.0, w, 2.62)
        icon_disc(g.shapes, x + 0.45, 3.42, 0.6, icon(ic), col)
        text(g.shapes, x + 0.85, 3.1, w - 1.0, 0.62, ttl, size=17, bold=True, color=NAVY, anchor="m")
        text(g.shapes, x + 0.25, 3.78, w - 0.45, 1.2, [{"t": t, "bullet": BLUE} for t in items], size=14,
             color=INK, space_after=2, line=1.0)
        if i == 1:
            text(g.shapes, x + 0.55, 4.72, w - 0.8, 0.42,
                 "~~*ω*_{b,p} = −*ω*_{b,x}·sin *q*_{y} + *ω*_{b,y}·cos *q*_{y}~~", size=18, color=NAVY, anchor="m")
        box(g.shapes, x + 0.2, 5.14, w - 0.4, 0.4, fill=ftint, radius=0.08, text=foot, size=13, bold=True,
            color=fcol)
        cards.append(g)
    c1 = chip(S, LM, 5.85, 6.0, 0.62, "Điểm kỳ dị ~~*q*_{p}~~ = ±90°: chặn |cos ~~*q*_{p}~~| ≥ cos 80°",
              fill="FFFFFF", color=NAVY, line_=BLUE, size=14, radius=0.31)
    c2 = chip(S, 6.8, 5.85, 6.0, 0.62, "IMU2 chỉ dùng sau: chờ yên → hiệu chuẩn → kiểm chứng", fill="FFFFFF",
              color=NAVY, line_=BLUE, size=14, radius=0.31)
    a = Anim(s)
    a.click()
    for i, c in enumerate(chev):
        a.add(c, "wipe_l", 300, after=i > 0)
    a.click().add(cards[0], "rise", 500)
    a.click().add(cards[1], "rise", 500)
    a.click().add(c1, "fade", 400).add(c2, "fade", 400, delay=200)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 11)


# =============================================================================
# 12. BO DIEU KHIEN NOI TANG (CASCADE)
# =============================================================================
def sum_node(S, cx, cy, d=0.36):
    return box(S, cx - d / 2, cy - d / 2, d, d, kind=MSO_SHAPE.OVAL, fill="FFFFFF", line=NAVY, lw=1.75,
               text="Σ", size=13, bold=True, color=NAVY, margin=0)


def s12():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 3 · THUẬT TOÁN ĐIỀU KHIỂN", "Bộ điều khiển nối tầng góc – tốc độ góc và vị trí các khâu bù")
    ym = 2.75
    blk = dict(radius=0.08, size=13, color="FFFFFF", line=0.95)
    # vong ngoai
    G1 = group(S)
    text(G1.shapes, LM, ym - 0.3, 0.6, 0.6, "~~*θ*_{ref}~~", size=18, color=NAVY, anchor="m")
    arrow(G1.shapes, 1.02, ym, 1.12, ym, color=NAVY)
    sum_node(G1.shapes, 1.3, ym)
    arrow(G1.shapes, 1.48, ym, 1.7, ym, color=NAVY)
    box(G1.shapes, 1.7, ym - 0.45, 1.85, 0.9, fill=BLUE,
        text=[{"t": "**Vòng góc · 100 Hz**"}, {"t": "P, giới hạn theo quãng đường dừng", "size": 11.5}], **blk)
    arrow(G1.shapes, 3.55, ym, 3.77, ym, color=NAVY)
    text(G1.shapes, 3.3, ym - 0.62, 0.75, 0.3, "~~*ω*_{sp}~~", size=15, color=NAVY, align="c")
    # vong trong
    G2 = group(S)
    sum_node(G2.shapes, 3.95, ym)
    arrow(G2.shapes, 4.13, ym, 4.35, ym, color=NAVY)
    box(G2.shapes, 4.35, ym - 0.45, 1.85, 0.9, fill=BLUE,
        text=[{"t": "**Vòng tốc độ 500 Hz**"}, {"t": "PI + truyền thẳng ~~*K*_{ff}~~", "size": 11.5}], **blk)
    # dong co + phan hoi
    G3 = group(S)
    box(G3.shapes, 11.6, ym - 0.45, 1.2, 0.9, fill=NAVY, text=[{"t": "**Động cơ**"}, {"t": "MS3506", "size": 12}],
        **blk)
    arrow(G3.shapes, 12.2, ym + 0.45, 12.2, 3.55, color=NAVY)
    box(G3.shapes, 10.25, 3.55, 2.55, 0.6, fill=TINT, radius=0.08, text="IMU1: ~~*θ*~~, ~~*ω*~~ camera", size=13,
        bold=True, color=NAVY)
    line(G3.shapes, 10.25, 3.72, 3.95, 3.72, color=YAW, w=1.75, tail=None)
    arrow(G3.shapes, 3.95, 3.72, 3.95, ym + 0.19, color=YAW, w=1.75)
    text(G3.shapes, 4.02, 3.36, 0.5, 0.3, "~~*ω*~~", size=15, color=YAW)
    line(G3.shapes, 10.25, 4.0, 1.3, 4.0, color=NAVY, w=1.75, tail=None)
    arrow(G3.shapes, 1.3, 4.0, 1.3, ym + 0.19, color=NAVY, w=1.75)
    text(G3.shapes, 1.37, 3.62, 0.5, 0.3, "~~*θ*~~", size=15, color=NAVY)
    text(G3.shapes, 1.02, ym + 0.12, 0.3, 0.3, "−", size=16, bold=True, color=BAD)
    text(G3.shapes, 3.67, ym + 0.12, 0.3, 0.3, "−", size=16, bold=True, color=BAD)
    # chuoi khau bo sung
    G4 = group(S)
    arrow(G4.shapes, 6.2, ym, 6.42, ym, color=NAVY)
    sum_node(G4.shapes, 6.6, ym)
    arrow(G4.shapes, 6.78, ym, 7.0, ym, color=NAVY)
    for x, w, t in ((7.0, 1.3, "④ Giới hạn\ntheo vùng"), (8.5, 1.55, "⑤ Đảo chiều,\ndừng, tạo dạng"),
                    (10.25, 1.15, "⑥ Lịch\nRS485")):
        box(G4.shapes, x, ym - 0.45, w, 0.9, fill="FFFFFF", line=ACCENT, lw=2.0, radius=0.08, text=t, size=13,
            bold=True, color=NAVY)
        arrow(G4.shapes, x + w, ym, x + w + (0.2 if x < 10 else 0.2), ym, color=NAVY)
    G5 = group(S)
    box(G5.shapes, 3.3, 1.38, 1.65, 0.5, fill=TINT, radius=0.08, text="IMU2: ~~*ω*_{b}~~, ~~*α*_{b}~~", size=13,
        bold=True, color=NAVY)
    arrow(G5.shapes, 4.95, 1.63, 5.15, 1.63, color=ACCENT, w=2.0)
    box(G5.shapes, 5.15, 1.38, 3.1, 0.5, fill=ACCENT, radius=0.08, text="② Bù IMU2 + ③ ngoại suy 12 ms", size=13,
        bold=True, color="FFFFFF")
    arrow(G5.shapes, 6.6, 1.88, 6.6, ym - 0.19, color=ACCENT, w=2.0)
    text(G5.shapes, 6.68, 2.0, 0.6, 0.3, "~~*u*_{ff}~~", size=15, color=ACCENT)
    # luat dieu khien + he so
    L = group(S)
    card(L.shapes, LM, 4.45, 6.85, 2.17)
    label(L.shapes, LM + 0.18, 4.52, 6.5, "LUẬT ĐIỀU KHIỂN RỜI RẠC (2.11), (2.12)", ico="code")
    text(L.shapes, LM + 0.2, 4.9, 6.5, 0.5,
         "~~*ω*_{sp,k} = sat( *K*_{p}^{θ}*e*_{θ,k} , min( *ω*_{max} , √(2*a*_{br}|*e*_{θ,k}|) ) )~~", size=16.5,
         color=NAVY, anchor="m")
    text(L.shapes, LM + 0.2, 5.4, 6.5, 0.5,
         "~~*u*_{k} = *K*_{ff}*ω*_{sp,k} + *K*_{p}^{ω}*ε*_{k} + *I*_{ω,k} + *u*_{ff,k} − *u*_{d,k}~~", size=16.5,
         color=NAVY, anchor="m")
    text(L.shapes, LM + 0.2, 5.92, 6.5, 0.66,
         "Vòng ngoài thuần P (~~*K*_{i}^{θ}~~ = 0): tránh hai tích phân lồng nhau · tích phân có điều kiện chống bão hòa · "
         "lệnh bù cộng **trước** khâu giới hạn", size=12.5, color=INK, line=0.98)
    T = group(S)
    card(T.shapes, 7.6, 4.45, 5.2, 2.17)
    rows = [["Hệ số", "● Yaw", "■ Pitch"],
            ["~~*K*_{p}^{θ}~~ (s^{−1})", "9,5", "38,0"], ["~~*K*_{p}^{ω}~~", "0,28", "0,15"],
            ["~~*K*_{i}^{ω}~~ (s^{−1})", "1,0", "0,5"], ["~~*ω*_{max}~~ (°/s)", "110", "100"]]
    table(T.shapes, 7.75, 4.56, [2.1, 1.45, 1.45], 0.39, rows, size=14.5)
    a = Anim(s)
    a.click().add(G1, "wipe_l", 500).add(G2, "wipe_l", 500, after=True).add(G3, "fade", 600, after=True)
    a.click().add(G4, "wipe_l", 600).add(G5, "wipe_t", 500, after=True)
    a.click().add(L, "rise", 500).add(T, "rise", 500, delay=200)
    a.build()
    set_transition(s, "push", dir="u")
    number(s)
    notes(s, 12)


# =============================================================================
# 13. BU IMU2 VA NGOAI SUY BU TRE
# =============================================================================
def s13():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 3 · THUẬT TOÁN ĐIỀU KHIỂN", "② Bù trước bằng IMU2 và ③ ngoại suy bù trễ 12 ms")
    card(S, LM, 1.38, 6.0, 5.24)
    label(S, LM + 0.18, 1.45, 5.6, "② BÙ TRUYỀN THẲNG TỪ IMU2", ico="exchange")
    p1 = group(S)
    pic(p1.shapes, icon("warning", BAD), 0.75, 1.9, 0.3, 0.3)
    text(p1.shapes, 1.15, 1.83, 5.2, 0.44, "Phản hồi chỉ tác động khi sai lệch đã xuất hiện", size=14.5, color=INK,
         anchor="m")
    dg = group(S)
    D = dg.shapes
    yb = 2.75
    box(D, 0.72, yb, 1.2, 0.6, fill=TINT, radius=0.08, text="IMU2 ~~*ω*_{b}~~", size=13, bold=True, color=NAVY)
    arrow(D, 1.92, yb + 0.3, 2.12, yb + 0.3, color=NAVY)
    box(D, 2.12, yb, 0.98, 0.6, fill=ACCENT, radius=0.08, text="× ~~*k*(*v*)~~", size=14, bold=True, color="FFFFFF")
    arrow(D, 3.1, yb + 0.3, 3.3, yb + 0.3, color=NAVY)
    sum_node(D, 3.48, yb + 0.3)
    text(D, 2.85, 2.28, 1.3, 0.28, "từ vòng góc", size=11.5, color=MUTED, align="c")
    arrow(D, 3.48, 2.56, 3.48, yb + 0.12, color=NAVY)
    arrow(D, 3.66, yb + 0.3, 3.86, yb + 0.3, color=NAVY)
    box(D, 3.86, yb, 1.4, 0.6, fill=BLUE, radius=0.08, text="Vòng tốc độ", size=13, bold=True, color="FFFFFF")
    arrow(D, 5.26, yb + 0.3, 5.44, yb + 0.3, color=NAVY)
    box(D, 5.44, yb, 0.95, 0.6, fill=NAVY, radius=0.08, text="Động cơ", size=12.5, bold=True, color="FFFFFF")
    fm = group(S)
    text(fm.shapes, 0.7, 3.55, 5.7, 0.48, "~~*u*_{ff} = *k*(*v*)·( *ω*_{b} + *τ*_{p}*α*_{b} )~~", size=19, color=NAVY,
         align="c", anchor="m")
    text(fm.shapes, 0.7, 4.03, 5.7, 0.48, "~~*e*_{θ,ss}^{ff} = −(1 − *k*)·*ω*_{b} / *K*_{p}^{θ}~~   (1.30)", size=19,
         color=NAVY, align="c", anchor="m")
    pill = box(S, 0.75, 4.58, 5.6, 0.7, fill=ACC_TINT, line=ACCENT, lw=1.5, radius=0.1,
               text="~~*k*~~ = 0,9 ⇒ sai lệch xác lập giảm **10 lần**, không tăng hệ số vòng kín", size=14.5,
               color=INK)
    nt = text(S, 0.75, 5.4, 5.65, 1.18, [
        {"t": "~~*k*(*v*)~~ tăng theo vùng tốc độ: 0,85 → 0,98 (Yaw)", "bullet": BLUE},
        {"t": "Cổng tin cậy: IMU2 ≤ 12 ms, giá trị hữu hạn; bù quá mức → chỉ giảm ~~*k*~~", "bullet": BLUE},
        {"t": "Không đổi hàm truyền hở → không lấn dự trữ ổn định", "bullet": BLUE}],
        size=13, color=INK, space_after=1, line=0.98)
    X = 6.75
    card(S, X, 1.38, 6.05, 5.24)
    label(S, X + 0.18, 1.45, 5.6, "③ NGOẠI SUY BÙ TRỄ", ico="stopwatch")
    p2 = group(S)
    pic(p2.shapes, icon("warning", BAD), X + 0.2, 1.9, 0.3, 0.3)
    text(p2.shapes, X + 0.6, 1.83, 5.3, 0.44, "Trễ ≈ 12 ms: lệnh bù đến muộn, đảo chiều thì ngược dấu", size=14,
         color=INK, anchor="m")
    cg = group(S)
    pic(cg.shapes, chart("s13_extrap"), X + 0.15, 2.32, 5.75, 2.75)
    chip(cg.shapes, X + 4.75, 4.72, 1.1, 0.3, "minh họa", fill="F1F4F8", color=MUTED, size=11, bold=False)
    f2 = group(S)
    text(f2.shapes, X + 0.15, 5.1, 5.75, 0.5, "~~*ω̂*_{b} = *ω*_{b} + sat( *τ*_{p}*α*_{b} , Δ*ω*_{max} )~~   (2.14)",
         size=18, color=NAVY, align="c", anchor="m")
    text(f2.shapes, X + 0.2, 5.62, 5.65, 0.95,
         "~~*τ*_{p}~~ = 12 ms · Δ~~*ω*_{max}~~ = 90 °/s · ~~*α*_{b}~~ lọc 18 Hz, vùng chết 70 °/s² · "
         "IMU2 quá 12 ms → giảm bù về 0 trong 30 ms", size=13, color=INK, align="c", line=1.0)
    a = Anim(s)
    a.click().add(p1, "fade", 400).add(dg, "wipe_l", 800, after=True)
    a.click().add(fm, "fade", 500).add(pill, "zoom", 450, after=True).add(pill, "pulse", 300, after=True)
    a.add(nt, "fade", 400, after=True)
    a.click().add(p2, "fade", 400).add(cg, "wipe_l", 900, after=True)
    a.click().add(f2, "rise", 500)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 13)


# =============================================================================
# 14. GIOI HAN LENH THEO TRANG THAI DONG
# =============================================================================
def s14():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 3 · THUẬT TOÁN ĐIỀU KHIỂN", "④ Giới hạn lệnh động: nới khi khung quay nhanh, chặt khi tĩnh")
    pb = group(S)
    card(pb.shapes, LM, 1.38, 4.35, 1.58)
    label(pb.shapes, LM + 0.18, 1.45, 4.0, "VẤN ĐỀ", ico="warning", color=BAD)
    text(pb.shapes, LM + 0.2, 1.8, 4.0, 1.12,
         "Giới hạn cố định thấp → bão hòa (**2,6 %** mẫu Yaw chạm 135 °/s); cao → lệnh lớn ngay cả khi tĩnh",
         size=14.5, color=INK, line=1.0)
    sol = group(S)
    card(sol.shapes, LM, 3.08, 4.35, 1.5)
    label(sol.shapes, LM + 0.18, 3.15, 4.0, "GIẢI PHÁP", ico="lightbulb", color=ACCENT)
    text(sol.shapes, LM + 0.2, 3.5, 4.0, 1.05,
         "4 vùng theo tốc độ khung ~~*v*~~; chuyển vùng có trễ; nội suy giới hạn theo ~~*λ*~~ (2.17)–(2.19) → "
         "lệnh không nhảy bậc", size=14, color=INK, line=1.0)
    rows = [["Vùng", "Ngưỡng vào\nYaw / Pitch", "~~*k*~~ Yaw / Pitch", "Giới hạn\n(°/s)"],
            ["Chậm", "–", "0,85 / 0,80", "110"], ["Thường", "50 / 30", "0,90 / 0,86", "135"],
            ["Nhanh", "140 / 90", "0,95 / 0,92", "165"], ["Rất nhanh", "220 / 170", "~~*k*(*v*)~~", "165"]]
    tb = table(S, LM, 4.72, [0.95, 1.2, 1.3, 0.9], 0.36, rows, size=12.5, row_h=[0.5, 0.35, 0.35, 0.35, 0.35])
    # bac thang
    X, W = 5.1, 7.7
    card(S, X, 1.38, W, 5.24)
    text(S, X + 0.2, 1.45, 6.5, 0.32, "BỐN VÙNG TRẠNG THÁI ĐỘNG · TRỤC YAW", size=12.5, bold=True, color=BLUE,
         spc=40)
    base_y = 5.72
    sc = 3.25 / 165
    zones = [("Chậm", 110, "0,85", "9CC3E8", INK), ("Thường", 135, "0,90", "5E9FD8", "FFFFFF"),
             ("Nhanh", 165, "0,95", "2F78C0", "FFFFFF"), ("Rất nhanh", 165, "k(v)", NAVY, "FFFFFF")]
    cols = []
    for i, (nm, lim, kk, fill, tc) in enumerate(zones):
        x = 5.75 + i * 1.7
        h = lim * sc
        g = group(S)
        box(g.shapes, x, base_y - h, 1.5, h, fill=fill, radius=0.06)
        text(g.shapes, x, base_y - h + 0.1, 1.5, 0.45, f"{lim} °/s", size=18, bold=True, color=tc, align="c")
        text(g.shapes, x, base_y - h + 0.55, 1.5, 0.35, "k = " + kk if kk != "k(v)" else "k = k(v)", size=13.5,
             color=tc, align="c", italic=kk == "k(v)")
        text(g.shapes, x, base_y - 0.48, 1.5, 0.4, nm, size=15, bold=True, color=tc, align="c")
        cols.append(g)
    axis = group(S)
    line(axis.shapes, 5.55, base_y, 12.6, base_y, color=MUTED, w=1.5, tail="tri")
    for i, v in enumerate(("v > 50", "v > 140", "v > 220")):
        x = 5.75 + (i + 1) * 1.7 - 0.1
        line(axis.shapes, x, base_y - 0.08, x, base_y + 0.12, color=MUTED, w=1.25, tail=None)
        text(axis.shapes, x - 0.6, base_y + 0.1, 1.2, 0.3, v, size=12.5, color=INK, align="c", bold=True)
    text(axis.shapes, 9.2, base_y + 0.4, 3.45, 0.3, "tốc độ khung mang ~~*v*~~ (°/s) →", size=12.5, color=MUTED,
         align="r")
    hy = group(S)
    box(hy.shapes, 5.5, 1.92, 3.1, 1.12, fill=ACC_TINT, radius=0.1, line=ACCENT, lw=1.25,
        text="Chuyển vùng có trễ: ngưỡng ra thấp hơn ngưỡng vào **25–30 %** → không dao động qua lại", size=13,
        color=INK, align="l", margin=0.1)
    a = Anim(s)
    a.click().add(pb, "rise", 450)
    a.click().add(axis, "wipe_l", 500)
    for c in cols:
        a.add(c, "wipe_b", 450, after=True)
    a.click().add(sol, "rise", 450).add(hy, "zoom", 450, delay=200)
    a.click()
    for i, r in enumerate(tb):
        a.add(r, "fade", 300, delay=0 if i == 0 else 100)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 14)


# =============================================================================
# 15. DAO CHIEU, DUNG, TAO DANG LENH
# =============================================================================
def flow_blocks(S, x, y0, w, items, h=0.66, gap=0.24, fills=None):
    out = []
    for i, t in enumerate(items):
        y = y0 + i * (h + gap)
        f = (fills[i] if fills else TINT)
        out.append(box(S, x, y, w, h, fill=f, radius=0.08, text=t, size=13.5, color=INK, line=0.98))
        if i < len(items) - 1:
            out.append(arrow(S, x + w / 2, y + h + 0.02, x + w / 2, y + h + gap - 0.02, color=MUTED, w=2.0))
    return out


def s15():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 3 · THUẬT TOÁN ĐIỀU KHIỂN",
           "⑤ Xử lý đảo chiều, dừng đột ngột và tạo dạng lệnh")
    w = (CW - 0.4) / 3
    colsh = []
    specs = [
        ("exchange", "ĐẢO CHIỀU", "Đổi dấu đột ngột → bước nhảy, rung; đổi chậm → camera bị khung kéo theo",
         ["Dự báo thời điểm qua 0:\n~~*t*_{0} = −*ω*_{b} / *α*_{b}~~", "~~*t*_{0}~~ ≤ 20 ms: mở cửa sổ,\ngiảm biên độ phía cũ",
          "Dấu mới đủ 2 chu kỳ\n→ chuyển phía"], None,
         "Dự báo để **phanh sớm**, đo được mới **đổi dấu**"),
        ("hand", "DỪNG ĐỘT NGỘT", "Bù và tích phân tồn dư kéo camera vượt ngược, phục hồi chậm",
         ["Chuyển động", "Phanh · 100 ms", "Ổn định · 120 ms", "Đã dừng: bù giảm về 0"],
         [TINT, TINT, TINT, ACC_TINT],
         "Xác nhận dừng: 3 điều kiện trong 220 ms · giữ ~~*T*_{h}~~ = 150 ms, giảm ~~*T*_{r}~~ = 50 ms"),
        ("sliders", "TẠO DẠNG LỆNH", "Lệnh nhảy bậc khi đảo chiều → dòng đỉnh lớn, rung cơ khí", None, None,
         "Giới hạn tốc độ biến thiên ~~*a*_{max}~~ = 18 000 °/s², nới 30 000 °/s² khi đảo chiều + giới hạn độ giật "
         "(2.20)–(2.22)"),
    ]
    for i, (ic, ttl, prob, items, fills, key) in enumerate(specs):
        x = LM + i * (w + 0.2)
        g = group(S)
        card(g.shapes, x, 1.38, w, 5.24)
        icon_disc(g.shapes, x + 0.42, 1.8, 0.56, icon(ic), BLUE)
        text(g.shapes, x + 0.8, 1.52, w - 0.9, 0.56, ttl, size=17, bold=True, color=NAVY, anchor="m", spc=40)
        text(g.shapes, x + 0.2, 2.2, w - 0.4, 0.72, prob, size=13.5, color=BAD, line=0.98)
        if items:
            hh = 0.66 if len(items) == 3 else 0.48
            gap = 0.2 if len(items) == 3 else 0.18
            flow_blocks(g.shapes, x + 0.25, 2.98, w - 0.5, items, h=hh, gap=gap, fills=fills)
        else:
            # minh hoa: lenh nhay bac (xam) va lenh tao dang (cam)
            ox, oy, ww, hh = x + 0.35, 3.05, w - 0.7, 1.95
            line(g.shapes, ox, oy + hh, ox + ww, oy + hh, color=MUTED, w=1.25, tail="tri")
            line(g.shapes, ox, oy + hh, ox, oy, color=MUTED, w=1.25, tail="tri")
            pts_g = [(0.0, 0.15), (0.3, 0.15), (0.3, 0.85), (1.0, 0.85)]
            for (a0, b0), (a1, b1) in zip(pts_g, pts_g[1:]):
                line(g.shapes, ox + a0 * ww, oy + b0 * hh, ox + a1 * ww, oy + b1 * hh, color="9AA7B4", w=2.5,
                     tail=None, dash=DASH)
            pts = [(0.0, 0.15), (0.3, 0.15)]
            yv = 0.15
            xv = 0.3
            for _ in range(5):
                yv2 = min(0.85, yv + 0.14)
                pts += [(xv, yv2), (xv + 0.09, yv2)]
                xv += 0.09
                yv = yv2
            pts += [(1.0, 0.85)]
            for (a0, b0), (a1, b1) in zip(pts, pts[1:]):
                line(g.shapes, ox + a0 * ww, oy + b0 * hh, ox + a1 * ww, oy + b1 * hh, color=ACCENT, w=2.75, tail=None)
            text(g.shapes, ox + 0.12, oy - 0.02, 0.6, 0.3, "~~*u*~~", size=13, color=MUTED)
            text(g.shapes, ox + ww - 0.4, oy + hh + 0.02, 0.4, 0.3, "~~*t*~~", size=13, color=MUTED, align="r")
            text(g.shapes, ox + 0.15, oy + hh * 0.52, 1.6, 0.3, "nhảy bậc", size=12, color=MUTED, bold=True)
            text(g.shapes, ox + 0.38 * ww + 0.1, oy + hh * 0.3, 1.9, 0.3, "tạo dạng: từng bước", size=12,
                 color=ACCENT, bold=True)
            chip(g.shapes, x + w - 1.3, 2.98, 1.05, 0.28, "minh họa", fill="F1F4F8", color=MUTED, size=10.5,
                 bold=False)
        box(g.shapes, x + 0.2, 5.55, w - 0.4, 0.92, fill=ACC_TINT, radius=0.1, text=key, size=13, color=INK,
            margin=0.08)
        colsh.append(g)
    a = Anim(s)
    for g in colsh:
        a.click().add(g, "rise", 500)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 15)


# =============================================================================
# 16. RS485 + GIAM SAT, BAO VE, KHOI PHUC
# =============================================================================
def s16():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 3 · TRUYỀN THÔNG VÀ GIÁM SÁT",
           "⑥ Lịch RS485 tất định và cơ chế giám sát – bảo vệ – khôi phục")
    card(S, LM, 1.38, 6.35, 5.24)
    label(S, LM + 0.18, 1.45, 6.0, "⑥ LỊCH RS485 LUÂN PHIÊN", ico="network")
    pb = group(S)
    pic(pb.shapes, icon("warning", BAD), 0.75, 1.9, 0.3, 0.3)
    text(pb.shapes, 1.15, 1.83, 5.6, 0.8,
         "2 động cơ chung bus bán song công; ưu tiên động → trễ ~~*τ*_{b}~~ ngẫu nhiên, phá giả thiết trễ hằng của ngoại suy",
         size=14.5, color=INK, line=1.0)
    x0, sw, yslot = 0.85, 0.72, 3.12
    slots = group(S)
    for i in range(8):
        nm, col = (("Yaw", YAW) if i % 2 == 0 else ("Pitch", PITCH))
        box(slots.shapes, x0 + i * sw + 0.02, yslot, sw - 0.04, 0.72, fill=col, radius=0.05, text=nm, size=13.5,
            bold=True, color="FFFFFF", margin=0)
    for i in range(0, 9, 2):
        text(slots.shapes, x0 + i * sw - 0.35, yslot + 0.76, 0.7, 0.26, f"{i * 5}" + (" ms" if i == 8 else ""),
             size=12, color=MUTED, align="c")
    tok = box(S, x0 + sw / 2 - 0.13, yslot - 0.36, 0.26, 0.26, kind=MSO_SHAPE.OVAL, fill=ACCENT)
    br = group(S)
    line(br.shapes, x0 + 0.02, 4.32, x0 + 2 * sw + 0.02, 4.32, color=NAVY, w=1.5, head="tri", tail="tri")
    text(br.shapes, x0 + 2 * sw + 0.12, 4.14, 3.9, 0.36, "mỗi trục cập nhật đều **10 ms**", size=15, color=NAVY,
         anchor="m")
    rules = text(S, 0.75, 4.72, 5.95, 1.85, [
        {"t": "Khe cố định 5 ms: khe chẵn Yaw, khe lẻ Pitch (2.28)", "bullet": BLUE},
        {"t": "Phát khung ở ưu tiên 116 → chờ phản hồi ≤ 3 ms", "bullet": BLUE},
        {"t": "**Không thử lại trong khe** (không lấn khe trục kia); 5 lỗi liên tiếp → báo sự cố", "bullet": BLUE}],
        size=15, color=INK, space_after=6, line=1.0)
    X = 7.1
    card(S, X, 1.38, 5.7, 5.24)
    label(S, X + 0.18, 1.45, 5.3, "GIÁM SÁT · BẢO VỆ · KHÔI PHỤC", ico="shield")
    nodes = {
        "KT": (7.35, 1.95, "Khởi tạo", TINT, NAVY), "SS": (10.4, 1.95, "Sẵn sàng", TINT, NAVY),
        "DK": (10.4, 3.25, "Điều khiển", BLUE, "FFFFFF"), "DAT": (7.35, 3.25, "Dừng an toàn", BAD_TINT, BAD),
    }
    nsh = {}
    for k, (x, y, t, f, c) in nodes.items():
        nsh[k] = box(S, x, y, 2.15, 0.6, fill=f, radius=0.1, text=t, size=14.5, bold=True, color=c,
                     line=BAD if k == "DAT" else None, lw=1.25)
    edges = group(S)
    arrow(edges.shapes, 9.5, 2.25, 10.4, 2.25, color=NAVY, w=1.75)
    arrow(edges.shapes, 11.15, 2.55, 11.15, 3.25, color=NAVY, w=1.75)
    text(edges.shapes, 10.38, 2.72, 0.75, 0.3, "START", size=11.5, bold=True, color=NAVY, align="r")
    arrow(edges.shapes, 11.75, 3.25, 11.75, 2.55, color=NAVY, w=1.75)
    text(edges.shapes, 11.82, 2.72, 0.75, 0.3, "STOP", size=11.5, bold=True, color=NAVY)
    arrow(edges.shapes, 10.4, 3.55, 9.5, 3.55, color=BAD, w=1.75)
    text(edges.shapes, 9.52, 3.6, 0.9, 0.3, "sự cố", size=11.5, bold=True, color=BAD, align="c")
    arrow(edges.shapes, 9.2, 3.25, 10.55, 2.55, color=GOOD, w=1.5)
    text(edges.shapes, 8.1, 2.62, 1.6, 0.3, "đã dừng hẳn", size=11.5, bold=True, color=GOOD, align="c")
    faults = text(S, X + 0.2, 4.18, 5.35, 2.4, [
        {"t": "Luồng giám sát an toàn (ưu tiên 118): dừng động cơ, ghi mã sự cố", "bullet": BAD},
        {"t": "Mất IMU1 hoặc tuổi mẫu > 6 ms → tắt điều khiển", "bullet": BAD},
        {"t": "Mất IMU2 → tắt bù; luồng khôi phục (108): chờ yên → hiệu chuẩn → kiểm chứng", "bullet": BAD},
        {"t": "RS485: 5 lỗi liên tiếp → báo sự cố", "bullet": BAD}],
        size=14.5, color=INK, space_after=6, line=1.0)
    a = Anim(s)
    a.click().add(pb, "fade", 400)
    a.click().add(slots, "wipe_l", 800).add(tok, "zoom", 300, after=True)
    a.add(tok, "path", 3600, after=True, path=f"M 0 0 L {7 * sw / 13.333:.4f} 0 E", ease=False)
    a.click().add(br, "fade", 400).add(rules, "fade", 450, delay=200)
    a.click().add(nsh["KT"], "zoom", 350).add(nsh["SS"], "zoom", 350, after=True)
    a.add(nsh["DK"], "zoom", 350, after=True).add(nsh["DAT"], "zoom", 350, after=True)
    a.add(edges, "fade", 500, after=True)
    a.click().add(faults, "fade", 500)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 16)


# =============================================================================
# 17. PHAN MEM GIAM SAT VA DIEU KHIEN TREN MAY TINH
# =============================================================================
def s17():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 3 · PHẦN MỀM GIÁM SÁT TRÊN MÁY TÍNH",
           "Phần mềm giám sát và điều khiển trên máy tính")
    gx, gy, gw = LM, 1.38, 8.85
    sc = gw / 1300.0
    gh = 741 * sc
    frame = box(S, gx - 0.04, gy - 0.04, gw + 0.08, gh + 0.08, fill="0B2545", radius=0.06, shd=True)
    shot = pic(S, img("gui"), gx, gy, gw, gh)
    groups = [
        ("Kết nối · START/STOP", "Cổng nối tiếp, tốc độ baud; bắt đầu / dừng điều khiển",
         [(286, 5, 557, 29), (1011, 423, 1287, 453)]),
        ("Điều khiển góc, tốc độ", "Đặt góc / tốc độ từng trục; chỉnh nhanh ±1°, ±10°, về 0°",
         [(1006, 455, 1291, 601)]),
        ("Theo dõi Yaw/Pitch, IMU", "Sai số từng trục, mô hình 3D, chỉ báo tư thế từ IMU",
         [(126, 39, 513, 85), (126, 511, 508, 699)]),
        ("Động cơ · RS485", "Dòng, tốc độ, nhiệt độ; tỉ lệ RS485; nhịp 500 / 100 Hz",
         [(1006, 91, 1293, 353)]),
        ("Đồ thị · nhật ký lỗi", "Đồ thị góc thời gian thực, chỉ tiêu; log sự kiện, lỗi",
         [(516, 91, 999, 616), (124, 707, 1297, 733)]),
    ]
    hl, labs = [], []
    X, W = 9.62, 3.18
    for i, (ttl, desc, rects) in enumerate(groups):
        g = group(S)
        for (x1, y1, x2, y2) in rects:
            box(g.shapes, gx + x1 * sc, gy + y1 * sc, (x2 - x1) * sc, (y2 - y1) * sc, fill=ACCENT, alpha=14,
                line=ACCENT, lw=2.75, radius=0.04)
        x1, y1 = rects[0][0], rects[0][1]
        bx = gx + x1 * sc - 0.16
        by = gy + y1 * sc - 0.16
        bx = max(gx - 0.12, bx)
        by = max(gy - 0.12, by)
        box(g.shapes, bx, by, 0.34, 0.34, kind=MSO_SHAPE.OVAL, fill=ACCENT, line="FFFFFF", lw=1.5, text=str(i + 1),
            size=12.5, bold=True, color="FFFFFF", margin=0)
        hl.append(g)
        y = 1.38 + i * 1.03
        lg = group(S)
        card(lg.shapes, X, y, W, 0.93)
        box(lg.shapes, X + 0.12, y + 0.2, 0.4, 0.4, kind=MSO_SHAPE.OVAL, fill=ACCENT, text=str(i + 1), size=14,
            bold=True, color="FFFFFF", margin=0)
        text(lg.shapes, X + 0.62, y + 0.05, W - 0.7, 0.36, ttl, size=13.5, bold=True, color=NAVY, anchor="m")
        text(lg.shapes, X + 0.62, y + 0.4, W - 0.7, 0.52, desc, size=11.5, color=INK, line=0.95)
        labs.append(lg)
    foot = text(S, X, 6.43, W, 0.3, "Python · PySide6 · ghi dữ liệu chu kỳ 10 ms", size=11.5, color=MUTED,
                italic=True)
    a = Anim(s)
    a.auto().add(frame, "fade", 400).add(shot, "fade", 400)
    for i in range(5):
        a.click()
        if i > 0:
            a.add(hl[i - 1], "exit_fade", 300)
            a.add(hl[i], "zoom", 400)
        else:
            a.add(hl[i], "zoom", 400)
        a.add(labs[i], "rise", 400, delay=100)
    a.click().add(hl[4], "exit_fade", 300).add(foot, "fade", 400)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 17)


# =============================================================================
# PHAN 4 - KHUON SLIDE HIEU CHINH (lo trinh + 5 khoi)
# =============================================================================
TRK_W, TRK_G = 1.272, 0.1


def tracker(s, cur):
    """Thanh lo trinh 9 buoc o dau slide; ten '!!trk_*' de Morph truot o sang."""
    S = s.shapes
    box(S, LM - 0.05, 0.16, CW + 0.1, 0.44, fill="FFFFFF", alpha=78, radius=0.12, name="!!trk_track")
    hx = LM + cur * (TRK_W + TRK_G)
    box(S, hx, 0.2, TRK_W, 0.36, fill=ACCENT, radius=0.1, name="!!trk_hl")
    for i, (short, _) in enumerate(STEPS):
        x = LM + i * (TRK_W + TRK_G)
        if i < cur:
            col, b, t = NAVY, False, "✓ " + short
        elif i == cur:
            col, b, t = "FFFFFF", True, short
        else:
            col, b, t = "8A99A8", False, short
        text(S, x, 0.2, TRK_W, 0.36, t, size=10.5, bold=b, color=col, align="c", anchor="m", name=f"!!trk_{i}")


def legend_ba(S, x, y):
    """Chu giai 'Truoc' (xam) / 'Sau' (mau truc)."""
    g = group(S)
    box(g.shapes, x, y + 0.08, 0.2, 0.2, fill=BEFORE, radius=0.03)
    text(g.shapes, x + 0.26, y, 0.75, 0.36, "Trước", size=12.5, color=INK, anchor="m")
    box(g.shapes, x + 1.0, y + 0.08, 0.13, 0.2, fill=YAW)
    box(g.shapes, x + 1.13, y + 0.08, 0.13, 0.2, fill=PITCH)
    text(g.shapes, x + 1.32, y, 0.6, 0.36, "Sau", size=12.5, color=INK, anchor="m")
    return g


def kpi_tile(S, x, y, w, h, spec):
    kind = spec[0]
    g = group(S)
    card(g.shapes, x, y, w, h)
    if kind == "single":
        _, val, lab, col = spec
        vs = 24 if len(val) <= 8 else (20 if len(val) <= 11 else 16)
        text(g.shapes, x + 0.06, y + 0.06, w - 0.12, 0.5, val, size=vs, bold=True, color=col, align="c", anchor="m")
        text(g.shapes, x + 0.08, y + 0.56, w - 0.16, h - 0.6, lab, size=12, color=INK, align="c", line=0.95)
    else:
        _, vy, vp, lab, col = spec
        if w > 2.6:
            text(g.shapes, x + 0.06, y + 0.06, w - 0.12, 0.5,
                 f"[[{YAW}:●]] {vy}    [[{PITCH}:■]] {vp}", size=19, bold=True, color=col, align="c", anchor="m")
            text(g.shapes, x + 0.08, y + 0.56, w - 0.16, h - 0.6, lab, size=12, color=INK, align="c", line=0.95)
        else:
            text(g.shapes, x + 0.06, y + 0.04, w - 0.12, 0.62,
                 [f"[[{YAW}:●]] {vy}", f"[[{PITCH}:■]] {vp}"], size=15, bold=True, color=col, align="c",
                 anchor="m", line=0.92)
            text(g.shapes, x + 0.08, y + 0.68, w - 0.16, h - 0.7, lab, size=11.5, color=INK, align="c", line=0.92)
    return g


def tune_slide(cur, title, prob, sim_label, sim_fn, choice, real_label, real_fn, kpis, key, ba_legend=True):
    s = new_slide()
    S = s.shapes
    tracker(s, cur)
    text(S, LM, 0.66, CW, 0.58, title, size=25, bold=True, color=NAVY, anchor="m", name="Title")
    P = group(S)
    card(P.shapes, LM, 1.38, 4.3, 1.12)
    label(P.shapes, LM + 0.15, 1.43, 4.0, "VẤN ĐỀ", color=BAD, ico="warning")
    text(P.shapes, LM + 0.18, 1.76, 3.98, 0.72, prob, size=13.5, color=INK, line=0.98)
    M = group(S)
    card(M.shapes, LM, 2.62, 4.3, 2.8)
    label(M.shapes, LM + 0.15, 2.67, 4.05, sim_label, ico="flask")
    sim_items = sim_fn(S, LM + 0.12, 3.05, 4.06, 2.3)
    C = group(S)
    box(C.shapes, LM, 5.54, 4.3, 1.08, fill=ACC_TINT, line=ACCENT, lw=1.5, radius=0.1)
    label(C.shapes, LM + 0.15, 5.57, 4.0, "LỰA CHỌN", color="B45309", ico="check")
    text(C.shapes, LM + 0.18, 5.88, 3.98, 0.72, choice, size=13.5, bold=True, color=INK, line=0.96)
    R = group(S)
    card(R.shapes, 5.05, 1.38, 7.75, 4.0)
    label(R.shapes, 5.2, 1.43, 5.6, real_label, ico="cogs", color=NAVY)
    lg = legend_ba(S, 10.82, 1.42) if ba_legend else None
    base_items, after_items = real_fn(S, 5.2, 1.86)
    tiles = []
    n = len(kpis)
    tw = (7.75 - 0.15 * (n - 1)) / n
    for i, sp in enumerate(kpis):
        tiles.append(kpi_tile(S, 5.05 + i * (tw + 0.15), 5.5, tw, 1.12, sp))
    a = Anim(s)
    a.auto().add(P, "fade", 400)
    a.click().add(M, "fade", 350)
    for it in sim_items:
        a.add(it, "wipe_l", 450, after=True)
    a.click().add(C, "zoom", 400).add(C, "pulse", 250, after=True)
    a.click().add(R, "fade", 350)
    if lg is not None:
        a.add(lg, "fade", 300)
    for it in base_items:
        a.add(it, "fade", 450, after=True)
    a.click()
    for i, it in enumerate(after_items):
        a.add(it, "wipe_l" if i == 0 else "zoom", 700 if i == 0 else 400, after=i > 0)
    a.click()
    for i, t in enumerate(tiles):
        a.add(t, "zoom", 380, delay=0 if i == 0 else 140)
    a.add(tiles[0], "pulse", 280, after=True)
    a.build()
    set_transition(s, "morph")
    number(s)
    notes(s, key)
    return s


def chart_layers(S, name, x, y):
    w, h = ANCH[name]["size"]
    files = ANCH[name]["files"]
    base = pic(S, chart(name, "base"), x, y, w, h)
    rest = [pic(S, chart(name, k), x, y, w, h) for k in files if k != "base"]
    return base, rest


def sim_chart(name):
    def f(S, x, y, w, h):
        cw, chh = ANCH[name]["size"]
        return [pic(S, chart(name), x, y, cw, chh)]
    return f


def sim_table(rows, widths, hl=None, row_h=None, size=13, extra=None):
    def f(S, x, y, w, h):
        out = table(S, x, y, widths, 0.42, rows, size=size, hl=hl, row_h=row_h)
        if extra:
            yy = y + sum(row_h) if row_h else y + 0.42 * len(rows)
            out.append(text(S, x, yy + 0.08, w, 0.4, extra, size=13, color=INK, anchor="m"))
        return out
    return f


def real_ba(name):
    def f(S, x, y):
        b, rest = chart_layers(S, name, x, y)
        return [b], rest
    return f


# =============================================================================
# 18. PHUONG PHAP: MO HINH MO PHONG + QUY TRINH HIEU CHINH
# =============================================================================
def s18():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 4 · MÔ PHỎNG, HIỆU CHỈNH VÀ THỰC NGHIỆM",
           "Mô phỏng thu hẹp miền tham số, hệ thật quyết định lựa chọn")
    text(S, LM, 1.36, 8, 0.3, "LỘ TRÌNH HOÀN THIỆN HỆ THỐNG · 7 BƯỚC HIỆU CHỈNH", size=12.5, bold=True,
         color=BLUE, spc=40)
    band = box(S, LM - 0.05, 1.7, CW + 0.1, 0.84, fill="FFFFFF", alpha=85, radius=0.12, name="!!trk_track", shd=True)
    chips = []
    for i, (short, full) in enumerate(STEPS):
        x = LM + i * (TRK_W + TRK_G)
        if i == 0:
            f, c = "8A99A8", "FFFFFF"
        elif i == 8:
            f, c = ACCENT, "FFFFFF"
        else:
            f, c = TINT, NAVY
        chips.append(box(S, x, 1.81, TRK_W, 0.62, fill=f, radius=0.1, text=full, size=12, bold=True, color=c,
                         line=0.92, margin=0.04, name=f"!!trk_{i}"))
    pr = group(S)
    card(pr.shapes, LM, 2.7, 4.2, 3.92)
    label(pr.shapes, LM + 0.15, 2.76, 4.0, "QUY TRÌNH MỖI BƯỚC", ico="route")
    rows = []
    for i, (ic, t) in enumerate((("warning", "Vấn đề hiện tại"), ("flask", "Mô phỏng khảo sát"),
                                 ("sliders", "Các phương án, tham số"), ("check", "Lựa chọn tham số"),
                                 ("cogs", "Áp dụng lên hệ thật"), ("balance", "So sánh trước / sau"),
                                 ("chartline", "Mức cải thiện · Yaw và Pitch"))):
        y = 3.13 + i * 0.43
        g = group(S)
        icon_disc(g.shapes, LM + 0.38, y + 0.19, 0.36, icon(ic), BAD if i == 0 else (ACCENT if i == 6 else BLUE),
                  icon_scale=0.55)
        text(g.shapes, LM + 0.68, y, 3.4, 0.38, t, size=14.5, color=INK, anchor="m", bold=i in (0, 6))
        rows.append(g)
    foot = text(S, LM + 0.15, 6.17, 3.95, 0.4, "Cấu hình sau mỗi bước = đối chứng của bước tiếp theo", size=12,
                color=MUTED, italic=True, anchor="m")
    X, W = 4.95, 7.85
    card(S, X, 2.7, W, 3.92)
    label(S, X + 0.15, 2.76, 4.0, "MÔ HÌNH MÔ PHỎNG (3.1)", ico="flask")
    fm = group(S)
    text(fm.shapes, X + 0.15, 3.12, 3.7, 0.52, "~~*G*_{m}(*s*) = exp(−*τ*_{d}*s*) / (*T*_{m}*s* + 1)~~", size=19,
         color=NAVY, anchor="m")
    text(fm.shapes, X + 3.95, 3.1, 3.8, 0.58, "~~*τ*_{d}~~ = 12 ms, bước tính 2 ms · ~~*T*_{m}~~ = 13 ms (Yaw), "
         "4 ms (Pitch)", size=12.5, color=INK, anchor="m", line=0.95)
    fig = pic(S, img("fig_model"), X + 0.17, 3.66, 7.5, 7.5 * 760 / 1920)
    ok = chip(S, X + 3.85, 2.76, 3.9, 0.34, "Sai lệch dạng sóng: Yaw 2,4 % · Pitch 2,2 %", fill=ACCENT, size=12)
    a = Anim(s)
    a.click().add(band, "fade", 300)
    for c in chips:
        a.add(c, "zoom", 260, after=True)
    a.click().add(pr, "fade", 300)
    for r in rows:
        a.add(r, "rise", 280, after=True)
    a.add(foot, "fade", 300, after=True)
    a.click().add(fm, "fade", 450)
    a.click().add(fig, "fade", 600).add(ok, "zoom", 400, after=True).add(ok, "pulse", 280, after=True)
    a.build()
    set_transition(s, "push", dir="u")
    number(s)
    notes(s, 18)


# =============================================================================
# 19 - 25. BAY BUOC HIEU CHINH
# =============================================================================
def s19():
    def real(S, x, y):
        fw = 4.0
        fh = fw * 757 / (1922 * (1 - 0.515))
        f = pic(S, img("fig_step"), x + 0.22, y + 0.06, fw, fh, crop=(0.515, 0, 0, 0))
        cap = text(S, x - 1.5 + 0.08, y + 0.06 + fh / 2 - 0.13, 3.0, 0.26, "đáp ứng chuẩn hóa (%)", size=11,
                   color=MUTED, italic=True, align="c", rot=270)
        rows = [["Hệ thật", "● Yaw", "■ Pitch"], ["~~*t*_{r}~~", "224,5 ms", "108,6 ms"],
                ["~~*t*_{s}~~", "520 ms", "265 ms"], ["Độ vọt lố", "0,398 %", "0,716 %"],
                ["σ giữ tĩnh", "0,0205°", "0,0068°"]]
        tb = table(S, x + 4.3, y + 0.35, [1.15, 1.07, 1.07], 0.56, rows, size=14)
        return [f, cap], tb

    tune_slide(
        1, "B1 · Hiệu chỉnh cascade: ~~*K*_{p}^{θ}~~ = 9,5 s^{−1} bám nhanh, còn dự trữ lệnh",
        "Cần bộ hệ số vừa bám nhanh, vọt lố nhỏ, vừa còn dự trữ lệnh cho các khâu bù",
        "MÔ PHỎNG · BẬC THANG TRỤC YAW",
        sim_table([["~~*K*_{p}^{θ}~~ (s^{−1})", "~~*t*_{r}~~ (ms)", "~~*t*_{s}~~ (ms)", "Lệnh / giới hạn"],
                   ["6", "299,6", "566", "35 %"], ["9,5", "162,7", "318", "56 %"], ["13", "97,1", "196", "76 %"]],
                  [1.06, 0.95, 0.95, 1.1], hl=2, row_h=[0.5, 0.42, 0.42, 0.42], size=13.5,
                  extra="~~*K*_{i}^{ω}~~ = 1,0 → ~~*J*_{e}~~ giảm **85,7 %**"),
        "Yaw: ~~*K*_{p}^{θ}~~ 9,5 · ~~*K*_{p}^{ω}~~ 0,28 · ~~*K*_{i}^{ω}~~ 1,0\n"
        "Pitch: 38,0 · 0,15 · 0,5 (quán tính nhỏ hơn)",
        "HỆ THẬT · BẬC THANG (HÌNH 3.7)", real,
        [("single", "−45,7 %", "~~*t*_{r}~~ so với ~~*K*_{p}^{θ}~~ = 6 (mô phỏng)", ACCENT),
         ("single", "−85,7 %", "~~*J*_{e}~~ nhờ ~~*K*_{i}^{ω}~~ = 1,0 (mô phỏng)", ACCENT),
         ("dual", "0,40 %", "0,72 %", "độ vọt lố < 1 % (hệ thật)", GOOD),
         ("dual", "20,24°", "13,24°", "~~*e*_{max}~~ đối chứng cho các bước sau", NAVY)],
        19, ba_legend=False)


def s20():
    tune_slide(
        2, "B2 · Bù IMU2: ~~*e*_{max}~~ giảm khoảng 10 % trên cả hai trục",
        "Phản hồi chỉ tác động khi sai lệch đã xuất hiện; sai lệch tỉ lệ tốc độ khung mang",
        "MÔ PHỎNG · ~~*e*_{rms}~~, KHUNG QUAY 240 °/s", sim_chart("s20_sim"),
        "Bù theo tốc độ: ~~*k*(*v*)~~ tăng theo vùng (Bảng 2.5), cộng vào đầu ra vòng tốc độ",
        "HỆ THẬT · BÀI THỬ ĐẢO CHIỀU NHANH 10,1 s", real_ba("s20_real"),
        [("single", "−10,3 %", "~~*e*_{max}~~ · Yaw", ACCENT), ("single", "−10,2 %", "~~*e*_{max}~~ · Pitch", ACCENT),
         ("dual", "−5,7 %", "−6,1 %", "~~*e*_{rms}~~ Yaw · Pitch", ACCENT),
         ("single", "+5,6 %", "~~*t*_{s}~~ khi khung nhanh (đánh đổi)", TRADE)], 20)


def s21():
    tune_slide(
        3, "B3 · Ngoại suy bù trễ ~~*τ*_{p}~~ = 12 ms: bước cải thiện lớn nhất",
        "Lệnh bù đến động cơ trễ ≈ 12 ms; khi khung đảo chiều, lượng bù còn ngược dấu",
        "MÔ PHỎNG · ~~*τ*_{p}~~ = 0 / 8 / 12 / 15 ms", sim_chart("s21_sim"),
        "~~*τ*_{p}~~ = 12 ms ≈ 80 % lợi ích của 15 ms, ít nhạy nhiễu ~~*α*_{b}~~; chặn 90 °/s",
        "HỆ THẬT · BÀI THỬ ĐẢO CHIỀU NHANH 10,1 s", real_ba("s21_real"),
        [("single", "−21,8 %", "~~*e*_{max}~~ · Yaw", ACCENT), ("single", "−23,5 %", "~~*e*_{max}~~ · Pitch", ACCENT),
         ("dual", "−17,2 %", "−17,5 %", "~~*e*_{rms}~~ Yaw · Pitch", ACCENT),
         ("single", "4,3 → 2,6 %", "tỉ lệ bão hòa ~~*ρ*_{sat}~~ Yaw", NAVY)], 21)


def s22():
    tune_slide(
        4, "B4 · Giới hạn lệnh động: tỉ lệ bão hòa Yaw giảm từ 2,6 % xuống 0,8 %",
        "Giới hạn cố định 135 °/s: 2,6 % mẫu Yaw bị bão hòa; nâng cố định → lệnh lớn khi tĩnh",
        "MÔ PHỎNG · ~~*e*_{rms}~~, XUNG 600 °/s", sim_chart("s22_sim"),
        "Luật 4 mốc 110 / 135 / 165 °/s; ngưỡng vào Yaw 50 / 140 / 220, Pitch 30 / 90 / 170 °/s",
        "HỆ THẬT · BÀI THỬ ĐẢO CHIỀU NHANH 10,1 s", real_ba("s22_real"),
        [("single", "−69 %", "~~*ρ*_{sat}~~ Yaw: 2,6 → 0,8 %", ACCENT), ("single", "−4,8 %", "~~*e*_{max}~~ · Yaw", ACCENT),
         ("single", "−0,2 %", "~~*e*_{max}~~ · Pitch", NAVY),
         ("single", "≈ 240 °/s", "giá thử thấp hơn 600 °/s của mô phỏng", MUTED)], 22)


def s23():
    def sim(S, x, y, w, h):
        g1 = group(S)
        box(g1.shapes, x, y, 1.9, 1.05, fill="F1F4F8", radius=0.08,
            text=[{"t": "**Ngắt bù tức thời**", "size": 13.5}, {"t": "bước nhảy lệnh, giật", "size": 12,
                                                                 "color": MUTED}], color=INK)
        g2 = group(S)
        box(g2.shapes, x + 2.06, y, 2.0, 1.05, fill=ACC_TINT, line=ACCENT, lw=1.75, radius=0.08,
            text=[{"t": "**Giảm bù dần** ★", "size": 13.5}, {"t": "giữ rồi giảm tuyến tính", "size": 12,
                                                              "color": MUTED}], color=INK)
        g3 = group(S)
        yy = y + 1.2
        for i, t in enumerate(("giữ ~~*T*_{h}~~ = 150 ms", "giảm ~~*T*_{r}~~ = 50 ms", "~~*ρ*_{I}~~ = 0,30",
                               "xác nhận dừng 220 ms", "cửa sổ ~~*t*_{0}~~ ≤ 20 ms", "xác nhận dấu 2 chu kỳ")):
            cx = x + (i % 2) * 2.06
            cy = yy + (i // 2) * 0.37
            chip(g3.shapes, cx, cy, 1.98, 0.31, t, fill=TINT, color=NAVY, size=12, bold=False)
        return [g1, g2, g3]

    def real(S, x, y):
        w = 6.95
        h = w * 682 / 1386
        f = pic(S, img("fig_stop"), x + 0.25, y + 0.02, w, h, crop=(0, 0.075, 0, 0))
        g = group(S)
        box(g.shapes, 8.55, 2.0, 1.55, 0.62, kind=MSO_SHAPE.OVAL, line=BAD, lw=2.5)
        chip(g.shapes, 10.12, 2.0, 2.3, 0.34, "không xử lý: vượt ngược", fill=BAD, size=11.5)
        return [f], [g]

    tune_slide(
        5, "B5 · Xử lý đảo chiều và dừng: thời gian phục hồi giảm hơn 50 %",
        "Khung dừng đột ngột: bù và tích phân tồn dư kéo camera vượt ngược, phục hồi chậm",
        "PHƯƠNG ÁN KHẢO SÁT (MÔ PHỎNG)", sim,
        "Giảm bù dần (giữ 150 ms, giảm 50 ms); xác nhận dừng 220 ms; đảo chiều xác nhận 2 chu kỳ",
        "HỆ THẬT · DỪNG ĐỘT NGỘT, 12 LẦN THỬ (HÌNH 3.17)", real,
        [("dual", "−52,4 %", "−58,7 %", "thời gian phục hồi sau dừng", ACCENT),
         ("single", "12/12", "xác nhận dừng (trước: 7/12, 8/12)", GOOD),
         ("dual", "−4,1 %", "−2,8 %", "~~*e*_{max}~~ khi đảo chiều", ACCENT),
         ("single", "≈ 2,40°", "sai lệch đỉnh không đổi (xảy ra lúc hãm)", NAVY)], 23, ba_legend=False)


def s24():
    def real(S, x, y):
        w = 7.2
        h = w * 632 / 1353
        f = pic(S, img("fig_shaping"), x + 0.12, y + 0.02, w, h, crop=(0, 0.088, 0, 0))
        g = group(S)
        chip(g.shapes, 6.3, 2.75, 1.3, 0.56, "nhảy 109 °/s\ntrong 2 ms", fill="7F8B98", size=11.5, radius=0.08)
        arrow(g.shapes, 7.6, 2.95, 7.88, 2.72, color="7F8B98", w=1.75)
        chip(g.shapes, 9.55, 3.22, 2.35, 0.36, "≤ 60 °/s mỗi chu kỳ 2 ms", fill=YAW, size=11.5, radius=0.08)
        arrow(g.shapes, 9.55, 3.4, 8.74, 3.52, color=YAW, w=1.75)
        return [f], [g]

    tune_slide(
        6, "B6 · Tạo dạng lệnh: bước lệnh giảm 45 %, dòng ~~*I*_{q}~~ giảm 22,3 %",
        "Khi đảo chiều, lệnh nhảy 109 °/s trong một chu kỳ 2 ms → dòng đỉnh, rung cơ khí",
        "MÔ PHỎNG · BẢNG 3.13",
        sim_table([["~~*a*_{max}~~ (°/s²)", "Bước lệnh (°/s)", "~~*e*_{max}~~ (°)", "Rung 10–30 Hz"],
                   ["Không giới hạn", "600", "19,87", "100 %"], ["18 000", "36", "21,40", "21 %"],
                   ["30 000", "60", "20,30", "38 %"], ["18 000 / 30 000", "60", "20,33", "23 %"]],
                  [1.42, 0.86, 0.82, 0.96], hl=4, row_h=[0.5, 0.4, 0.4, 0.4, 0.4], size=12.5),
        "18 000 °/s², nới 30 000 °/s² khi đảo chiều + giới hạn độ giật: rung còn 23 %, ~~*e*_{max}~~ +2,3 %",
        "HỆ THẬT · LỆNH TRỤC YAW KHI ĐẢO CHIỀU (HÌNH 3.18)", real,
        [("single", "−45,0 %", "bước lệnh lớn nhất 109,0 → 59,9 °/s", ACCENT),
         ("single", "−22,3 %", "dòng ~~*I*_{q}~~ hiệu dụng", ACCENT),
         ("single", "6 → 12 ms", "thời gian đổi dấu lệnh", NAVY),
         ("dual", "+1,4 %", "+1,1 %", "~~*e*_{rms}~~ (đánh đổi có chủ đích)", TRADE)], 24, ba_legend=False)


def s25():
    tune_slide(
        7, "B7 · Lịch RS485 luân phiên: Pitch cập nhật đều, độ tản mát giảm 85 %",
        "Ưu tiên Yaw trên bus chung: Pitch cập nhật không đều → trễ thay đổi, ngoại suy sai",
        "MÔ PHỎNG · 20 000 KHE",
        sim_table([["Cập nhật (ms)", "Ưu tiên Yaw", "Luân phiên"],
                   ["Pitch p99 / lớn nhất", "30,2 / 70,0", "11,8 / 22,2"],
                   ["Yaw p99 / lớn nhất", "7,1 / 15,1", "11,7 / 30,1"],
                   ["Không khởi động được", "7,2 %", "0,8 %"]],
                  [1.7, 1.18, 1.18], row_h=[0.46, 0.55, 0.55, 0.55], size=13),
        "Luân phiên, khe 5 ms: Yaw khe chẵn, Pitch khe lẻ; không thử lại trong khe",
        "HỆ THẬT · TẢI ĐẦY ĐỦ (BẢNG 3.19)", real_ba("s25_real"),
        [("single", "−85 %", "tản mát cập nhật Pitch 13,3 → 2,0 ms", ACCENT),
         ("single", "62,1 → 14,1", "cập nhật lớn nhất Pitch (ms)", ACCENT),
         ("single", "−18,5 %", "~~*e*_{max}~~ Pitch so với ưu tiên Yaw", ACCENT),
         ("single", "6,00 ms", "giao dịch lớn nhất > khe 5 ms (chưa đạt)", BAD)], 25, ba_legend=False)


# =============================================================================
# 26. TONG HOP MUC CAI THIEN QUA TUNG BUOC
# =============================================================================
def s26():
    s = new_slide()
    S = s.shapes
    tracker(s, 8)
    text(S, LM, 0.66, CW, 0.58, "Hoàn thiện dần qua 7 bước: ~~*e*_{max}~~ giảm 34,9 % (Yaw) và 32,6 % (Pitch)",
         size=25, bold=True, color=NAVY, anchor="m", name="Title")
    card(S, LM, 1.38, 8.55, 5.24)
    base, segs = chart_layers(S, "s26_chain", LM + 0.15, 1.5)
    X, W = 9.3, 3.5
    specs = [("dual", "−34,9 %", "−32,6 %", "~~*e*_{max}~~ sau 7 bước", ACCENT),
             ("dual", "−23,3 %", "−22,8 %", "~~*e*_{rms}~~ sau 7 bước", ACCENT),
             ("single", "6,1 → 0,7 %", "tỉ lệ bão hòa ~~*ρ*_{sat}~~ Yaw (−88,5 %)", ACCENT),
             ("single", "−23,7 %", "phục hồi sau đảo chiều 0,392 → 0,299 s", ACCENT),
             ("dual", "−52,4 %", "−58,7 %", "phục hồi sau dừng", ACCENT)]
    tiles = [kpi_tile(S, X, 1.38 + i * 1.07, W, 0.97, sp) for i, sp in enumerate(specs)]
    a = Anim(s)
    a.click().add(segs[0], "wipe_l", 700)
    a.click().add(segs[1], "wipe_l", 700)
    a.click().add(segs[2], "wipe_l", 600).add(segs[3], "wipe_l", 600, after=True)
    a.click().add(segs[4], "wipe_l", 600).add(segs[5], "wipe_l", 600, after=True)
    a.click()
    for i, t in enumerate(tiles):
        a.add(t, "rise", 380, delay=0 if i == 0 else 130)
    a.add(tiles[0], "pulse", 300, after=True)
    a.build()
    set_transition(s, "morph")
    number(s)
    notes(s, 26)


# =============================================================================
# 27. DANH GIA KET QUA THEO CHI TIEU
# =============================================================================
def s27():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 5 · ĐÁNH GIÁ VÀ KẾT LUẬN", "Đánh giá kết quả theo hệ chỉ tiêu kỹ thuật")
    cols = [
        ("checkcircle", GOOD, "ĐẠT", [
            "σ giữ tĩnh **0,0205° / 0,0068°** < 0,03°", "Độ vọt lố **0,40 % / 0,72 %** < 1 %",
            "Vòng tốc độ góc giữ đúng **2 ms**", "Xác nhận dừng **12/12**; 0 giao dịch lỗi"]),
        ("chartline", BLUE, "CẢI THIỆN", [
            "~~*e*_{rms}~~ **−23,3 % / −22,8 %**", "~~*e*_{max}~~ **−34,9 % / −32,6 %**",
            "Phục hồi sau dừng **−52,4 % / −58,7 %**", "Tản mát cập nhật Pitch **−85 %**",
            "Bước lệnh **−45 %**, dòng ~~*I*_{q}~~ **−22,3 %**"]),
        ("balance", TRADE, "ĐÁNH ĐỔI", [
            "~~*t*_{s}~~ **+5,6 %** khi khung nhanh (bù IMU2)", "~~*e*_{rms}~~ Yaw **+1,4 %** (tạo dạng lệnh)",
            "Không chọn ưu tiên Yaw dù ~~*e*_{rms}~~ Yaw thấp hơn 8,7 %"]),
        ("timescircle", BAD, "CHƯA ĐẠT", [
            "Giao dịch lớn nhất **6,00 ms** > khe 5 ms", "Cập nhật lớn nhất **14,1 ms** > 10 ms",
            "p99 giao dịch **+88,9 %** khi tải tăng"]),
    ]
    w = (CW - 3 * 0.18) / 4
    gs = []
    for i, (ic, col, ttl, items) in enumerate(cols):
        x = LM + i * (w + 0.18)
        g = group(S)
        card(g.shapes, x, 1.38, w, 4.62)
        icon_disc(g.shapes, x + w / 2, 1.98, 0.8, icon(ic), col)
        text(g.shapes, x + 0.1, 2.48, w - 0.2, 0.5, ttl, size=17, bold=True, color=col, align="c", anchor="m")
        text(g.shapes, x + 0.2, 3.05, w - 0.33, 2.9, [{"t": t, "bullet": col} for t in items], size=14.5,
             color=INK, space_after=6, line=1.0)
        gs.append(g)
    sc = box(S, LM, 6.15, CW, 0.47, fill=NAVY, radius=0.1,
             text="Phạm vi kiểm chứng: tốc độ camera hiệu dụng ≈ 50 °/s, đỉnh ≈ 240 °/s · hãm 2 300 °/s² · "
                  "bậc 2–7° (Yaw), 3° (Pitch)", size=13, color="FFFFFF")
    a = Anim(s)
    for i, g in enumerate(gs):
        a.click().add(g, "rise", 450)
        if i == 3:
            a.add(sc, "fade", 400, after=True)
    a.build()
    set_transition(s, "push", dir="u")
    number(s)
    notes(s, 27)


# =============================================================================
# 28. TONG HOP: CAC THANH PHAN -> HE GIMBAL HOAN THIEN
# =============================================================================
def s28():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 5 · TỔNG HỢP", "Các nội dung đã thực hiện tạo nên hệ Gimbal hoàn thiện")
    items = [("cube", "Thiết kế cơ khí", "thiết kế · chế tạo"), ("chip", "Thiết kế mạch", "thiết kế · chế tạo"),
             ("code", "Lập trình NuttX", "7 luồng thời gian thực"), ("compass", "Xử lý hai IMU", "lọc · ước lượng tư thế"),
             ("sitemap", "Điều khiển cascade", "góc – tốc độ góc"), ("exchange", "Bù IMU2", "bù trước khung mang"),
             ("stopwatch", "Bù trễ", "ngoại suy 12 ms"), ("route", "Xử lý chuyển động nhanh", "giới hạn · đảo chiều · dừng"),
             ("network", "RS485", "lịch luân phiên 5 ms"), ("desktop", "Phần mềm giám sát", "điều khiển · ghi log")]
    tw, th, gap = 1.5, 2.08, 0.27
    tiles, plus = [], []
    for i, (ic, t1, t2) in enumerate(items):
        r, c = divmod(i, 5)
        x = LM + c * (tw + gap)
        y = 1.45 + r * (th + 0.32)
        g = group(S)
        card(g.shapes, x, y, tw, th)
        icon_disc(g.shapes, x + tw / 2, y + 0.52, 0.66, icon(ic), BLUE if i < 4 else (ACCENT if i < 9 else NAVY))
        text(g.shapes, x + 0.06, y + 0.92, tw - 0.12, 0.66, t1, size=14, bold=True, color=NAVY, align="c", anchor="m",
             line=0.95)
        text(g.shapes, x + 0.06, y + 1.56, tw - 0.12, 0.46, t2, size=11.5, color=MUTED, align="c", line=0.95)
        tiles.append(g)
        if c < 4:
            plus.append(text(S, x + tw, y + th / 2 - 0.2, gap, 0.4, "+", size=20, bold=True, color=ACCENT,
                             align="c", anchor="m"))
        else:
            plus.append(None)
    arr = box(S, 9.15, 3.15, 0.6, 0.7, fill=ACCENT, kind=MSO_SHAPE.RIGHT_ARROW)
    fin = group(S)
    card(fin.shapes, 9.9, 1.45, 2.9, 4.5, fill="FFFFFF")
    pic(fin.shapes, img("cad_render_cut"), 10.0, 1.6, 2.7)
    text(fin.shapes, 9.95, 4.2, 2.8, 0.6, "HỆ GIMBAL\nHOÀN THIỆN", size=17, bold=True, color=NAVY, align="c",
         anchor="m", line=0.95)
    text(fin.shapes, 9.98, 4.85, 2.74, 1.0, [f"[[{YAW}:●]] ~~*e*_{{max}}~~ −34,9 %", f"[[{PITCH}:■]] ~~*e*_{{max}}~~ −32,6 %",
                                             "phục hồi sau dừng −52 %…−59 %"], size=12.5, color=INK, align="c",
         line=1.0)
    band = box(S, LM, 6.32, CW, 0.32, fill=None, text="Đã thiết kế · đã chế tạo · đã lập trình · đã mô phỏng · "
               "đã kiểm chứng trên hệ thật", size=13.5, bold=True, color=NAVY)
    a = Anim(s)
    a.click()
    for i, g in enumerate(tiles):
        a.add(g, "zoom", 280, after=i > 0)
        if plus[i] is not None:
            a.add(plus[i], "fade", 150)
    a.click().add(arr, "wipe_l", 400).add(fin, "zoom", 600, after=True).add(fin, "pulse", 300, after=True)
    a.add(band, "fade", 400, after=True)
    a.build()
    set_transition(s, "zoom")
    number(s)
    notes(s, 28)


# =============================================================================
# 29. KET LUAN, HAN CHE, HUONG PHAT TRIEN
# =============================================================================
def s29():
    s = new_slide()
    S = s.shapes
    header(s, "PHẦN 5 · KẾT LUẬN", "Kết luận, hạn chế và hướng phát triển")
    cols = [
        ("checkcircle", NAVY, "KẾT QUẢ ĐẠT ĐƯỢC", [
            "Thiết kế, chế tạo Gimbal 2 trục Yaw–Pitch, mạch STM32F407 và phần mềm giám sát",
            "Phần mềm Apache NuttX 7 luồng; xử lý 2 IMU; 6 khâu điều khiển",
            "Mô hình mô phỏng sai lệch dạng sóng 2,4 % / 2,2 %",
            "7 bước hiệu chỉnh: ~~*e*_{max}~~ −34,9 % / −32,6 %; ~~*e*_{rms}~~ −23,3 % / −22,8 %; "
            "phục hồi sau dừng −52,4 % / −58,7 %"]),
        ("warning", BAD, "HẠN CHẾ", [
            "RS485 chưa tất định hoàn toàn: giao dịch 6,00 ms > khe 5 ms; cập nhật lớn nhất 14,1 ms",
            "Thực nghiệm đến ≈ 240 °/s; dải 600–900 °/s mới xác nhận bằng mô phỏng",
            "Mô hình chưa có ma sát, khe hở, đàn hồi → lạc quan hơn hệ thật"]),
        ("route", ACCENT, "HƯỚNG PHÁT TRIỂN", [
            "Bus riêng mỗi trục hoặc CAN; ngoại suy thích nghi theo tuổi dữ liệu",
            "Quét tần số, xác nhận mode cộng hưởng, lọc chắn dải 16,5 Hz",
            "Giá thử kích thích > 240 °/s, lặp đủ số lần",
            "Tham chiếu phương vị độc lập cho trục Yaw",
            "Điều khiển trực tiếp dòng ~~*I*_{q}~~"]),
    ]
    w = (CW - 2 * 0.2) / 3
    gs = []
    for i, (ic, col, ttl, items) in enumerate(cols):
        x = LM + i * (w + 0.2)
        g = group(S)
        card(g.shapes, x, 1.38, w, 5.24)
        box(g.shapes, x, 1.38, w, 0.78, fill=col, radius=0.1)
        icon_disc(g.shapes, x + 0.5, 1.77, 0.54, icon(ic), col, ring="FFFFFF")
        text(g.shapes, x + 0.9, 1.38, w - 1.0, 0.78, ttl, size=17, bold=True, color="FFFFFF", anchor="m", spc=30)
        text(g.shapes, x + 0.22, 2.38, w - 0.4, 4.15, [{"t": t, "bullet": col} for t in items], size=16,
             color=INK, space_after=10, line=1.02)
        gs.append(g)
    a = Anim(s)
    for g in gs:
        a.click().add(g, "rise", 500)
    a.build()
    set_transition(s, "fade")
    number(s)
    notes(s, 29)


# =============================================================================
# 30. CAM ON
# =============================================================================
def s30():
    s = new_slide(L_WAVE)
    S = s.shapes
    pic(S, img("logo"), 6.16, 0.55, 1.0, 1.0)
    t1 = text(S, 0.8, 1.85, 11.73, 1.0, "EM XIN CHÂN THÀNH CẢM ƠN!", size=44, bold=True, color=NAVY, align="c",
              anchor="m")
    t2 = text(S, 0.8, 2.9, 11.73, 0.6, "Kính mong nhận được ý kiến đóng góp của Hội đồng", size=24, italic=True,
              color=BLUE, align="c", anchor="m")
    hero = pic(S, img("cad_render_cut"), 5.55, 3.55, 2.2)
    info = box(S, 1.2, 5.75, 10.93, 0.85, fill=NAVY, radius=0.12, text=[
        {"t": "Đề tài: Thiết kế bộ điều khiển Gimbal trên cơ sở hệ điều hành thời gian thực NuttX", "size": 15,
         "bold": True},
        {"t": "GVHD: Trung tá, TS. Nguyễn Ngọc Hưng · Trung tá, ThS. Lê Văn Huy   |   HV: Huỳnh Thanh Nam – "
              "Lớp Tên lửa phòng không 2", "size": 13}], color="FFFFFF")
    a = Anim(s)
    a.auto().add(t1, "zoom", 600).add(t2, "fade", 500, after=True).add(hero, "fade", 500, delay=100)
    a.add(info, "rise", 500, after=True)
    a.build()
    set_transition(s, "fade", "slow")
    number(s)
    notes(s, 30)


# =============================================================================
if __name__ == "__main__":
    for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14, s15, s16, s17, s18, s19,
               s20, s21, s22, s23, s24, s25, s26, s27, s28, s29, s30):
        fn()
    prs.save(OUT)
    print("saved", OUT, "slides:", len(prs.slides))
