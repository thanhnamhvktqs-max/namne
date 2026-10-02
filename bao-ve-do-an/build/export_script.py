# -*- coding: utf-8 -*-
"""Chen loi thoai (Notes) cua tung slide vao KICH_BAN_THUYET_TRINH.md.

    python export_script.py <deck.pptx> ../KICH_BAN_THUYET_TRINH.md
"""
import re
import sys

from pptx import Presentation

from notes import NOTES

deck, md = sys.argv[1], sys.argv[2]
prs = Presentation(deck)
parts = []
for i, s in enumerate(prs.slides, 1):
    title = next((sh.text_frame.text for sh in s.shapes if sh.name == "Title" and sh.has_text_frame), "")
    if i == 1:
        title = "Trang bìa"
    elif i == 30:
        title = "Cảm ơn"
    note = NOTES[i]
    m = re.match(r"\[(\d+:\d+)\]\s*", note)
    t = m.group(1) if m else ""
    body = note[m.end():] if m else note
    body = body.replace("(Nhấp)", "**(Nhấp)**")
    parts.append(f"### Slide {i} · {title.strip()}  _(≈ {t})_\n\n{body}\n")
text = open(md, encoding="utf-8").read()
start, end = "<!-- NOTES-START -->", "<!-- NOTES-END -->"
a, b = text.index(start) + len(start), text.index(end)
text = text[:a] + "\n\n" + "\n".join(parts) + "\n" + text[b:]
open(md, "w", encoding="utf-8").write(text)
print("ok", len(parts))
